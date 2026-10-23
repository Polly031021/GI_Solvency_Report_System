from __future__ import annotations

"""Regression tests for the STEP7 metric catalog, STEP3 gap check and unit default."""

import io
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from services.batch_extraction import FlexibleTable
from services.solvency_normalizer import STANDARD_COLUMNS
from services.solvency_step3_standardizer import (
    standardize_to_target,
    step3_metric_catalog,
)
from services.solvency_step7_metric_catalog import (
    load_taxonomy_workbook,
    step7_metric_requirement_catalog,
    step7_module_metric_rows,
)
from services.step3_metric_gap import (
    MISSING_SHEET_NAME,
    build_step7_metric_gap,
    merge_missing_metric_fills,
    missing_metric_template_workbook_bytes,
    read_missing_metric_fills,
)


def _row(code: str, name: str, value: float, *, unit: str = "万元", period: str = "本季度末数", data_type: str = "金额") -> dict:
    record = {column: "" for column in STANDARD_COLUMNS}
    record.update(
        {
            "公司": "测试财险",
            "原始公司名称": "测试财险",
            "标准公司名称": "测试财险",
            "公司统一编码": "C001",
            "公司类型": "财险",
            "同业分类": "中型公司",
            "报告类型": "NON_LIFE_SOLVENCY",
            "报告年度": "2025",
            "报告季度": "Q4",
            "报告期": "2025Q4",
            "指标编码": code,
            "指标名称": name,
            "期间口径": period,
            "数值": value,
            "单位": unit,
            "数据类型": data_type,
            "指标属性": "披露",
            "来源类型": "报告提取",
            "来源文件": "测试财险.pdf",
            "来源工作表": "S01",
            "导入批次": "测试财险:2025Q4",
        }
    )
    return record


def test_step7_catalog_covers_template_and_implicit_metrics() -> None:
    taxonomy = load_taxonomy_workbook()
    catalog = step7_metric_requirement_catalog(taxonomy)
    mapping = step7_module_metric_rows()
    codes = set(catalog["指标编码"].astype(str))

    assert len(catalog) >= 30, len(catalog)
    assert len(mapping) >= len(catalog), (len(mapping), len(catalog))
    assert catalog["指标编码"].is_unique
    for code in (
        "SIGNED_PREMIUM",
        "SIGNED_PREMIUM_MOTOR",
        "COMBINED_SOLVENCY_RATIO",
        "CORE_SOLVENCY_RATIO",
        "ACTUAL_CAPITAL",
        "MINIMUM_CAPITAL",
        "MARKET_RISK_TO_QUANT_CAPITAL",
        "NON_LIFE_CATASTROPHE_RISK_CAPITAL",
        "NON_LIFE_INSURANCE_RISK_DIVERSIFICATION_EFFECT",
        "LOSS_RATIO",
        "EXPENSE_RATIO",
        "INVESTMENT_RETURN",
    ):
        assert code in codes, code
    assert (catalog["补录方式"] == "直接补录").sum() >= 25
    assert (catalog["是否派生"] == "是").any()
    print("PASS STEP7 指标需求目录覆盖模板与代码隐含指标")


def test_missing_gap_template_and_merge() -> None:
    taxonomy = load_taxonomy_workbook()
    data = pd.DataFrame(
        [_row("SIGNED_PREMIUM", "签单保费", 100.0)],
        columns=STANDARD_COLUMNS,
    )
    report = build_step7_metric_gap(data, taxonomy=taxonomy)
    assert report.missing_count > 0
    assert "ACTUAL_CAPITAL" in set(report.missing["指标编码"])
    assert "SIGNED_PREMIUM" not in set(report.missing["指标编码"])
    assert report.coverage_ratio < 1

    workbook = missing_metric_template_workbook_bytes(report)
    wb = load_workbook(io.BytesIO(workbook))
    assert MISSING_SHEET_NAME in wb.sheetnames
    ws = wb[MISSING_SHEET_NAME]
    headers = [cell.value for cell in ws[1]]
    code_col = headers.index("指标编码") + 1
    value_col = headers.index("数值") + 1
    unit_col = headers.index("单位") + 1
    filled = False
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=code_col).value == "ACTUAL_CAPITAL":
            ws.cell(row=row, column=value_col, value=10000)
            ws.cell(row=row, column=unit_col, value="元")
            filled = True
            break
    assert filled
    output = io.BytesIO()
    wb.save(output)

    parsed = read_missing_metric_fills(output.getvalue())
    assert parsed.filled_count == 1
    assert parsed.skipped.empty, parsed.skipped
    merged = merge_missing_metric_fills(
        data,
        parsed.rows,
        taxonomy=taxonomy,
        requirement_catalog=report.requirement_catalog,
    )
    assert merged.added_count == 1
    actual = merged.data[merged.data["指标编码"].eq("ACTUAL_CAPITAL")].iloc[0]
    assert actual["数值"] == 1
    assert actual["单位"] == "万元"
    assert actual["来源类型"] == "人工补录"
    assert merged.data["指标编码"].eq("ACTUAL_CAPITAL").sum() == 1
    print("PASS STEP7 缺失指标导出、填写与合并")


def test_missing_gap_includes_empty_batch_company() -> None:
    taxonomy = load_taxonomy_workbook()
    data = pd.DataFrame(columns=STANDARD_COLUMNS)
    report = build_step7_metric_gap(
        data,
        taxonomy=taxonomy,
        company_periods=[
            {
                "公司": "空白财险",
                "报告期": "2025Q4",
                "来源文件": "空白财险.pdf",
            }
        ],
    )
    assert report.missing_count == report.required_base_count
    assert set(report.missing["公司"]) == {"空白财险"}
    print("PASS 无提取结果的公司也会显示 STEP7 缺失指标")


def test_step3_defaults_missing_amount_unit_to_yuan() -> None:
    table = FlexibleTable(
        table_id="ACTUAL_CAPITAL",
        table_name="S02-实际资本明细表",
        page=1,
        rows=[
            ["行次", "项目", "本季度末数", "上季度末数"],
            ["1", "实际资本", "10,000", "20,000"],
        ],
        unit="",
        unit_source="未识别",
    )
    taxonomy = load_taxonomy_workbook()
    catalog = step3_metric_catalog(taxonomy, include_derived=False)
    result = standardize_to_target(
        [table],
        taxonomy,
        {"公司": "测试财险", "报告期": "2025Q4", "来源文件": "测试财险.pdf"},
        "财险",
        catalog,
        report_profile_id="NON_LIFE_SOLVENCY",
        allowed_company_types=("财险",),
        include_derived=False,
    )
    row = result.data[result.data["指标编码"].eq("ACTUAL_CAPITAL")].iloc[0]
    assert row["数值"] == 1
    assert row["单位"] == "万元"
    assert "按默认单位元换算" in str(row["备注"])
    assert any("按默认单位" in warning for warning in result.warnings)
    print("PASS 未识别单位的金额默认按元换算为万元")


if __name__ == "__main__":
    test_step7_catalog_covers_template_and_implicit_metrics()
    test_missing_gap_template_and_merge()
    test_missing_gap_includes_empty_batch_company()
    test_step3_defaults_missing_amount_unit_to_yuan()
    print("全部通过")
