# -*- coding: utf-8 -*-
"""STEP1-STEP4 多文件文件池、规则提取与单工作表合并回归测试。"""

from __future__ import annotations

import io
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
from streamlit.testing.v1 import AppTest

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from services.batch_processor import build_context_from_profile  # noqa: E402
from services.batch_workflow import (  # noqa: E402
    document_pdf_bytes,
    extract_document_tables,
    merge_standard_frames,
    single_sheet_workbook_bytes,
    standardize_document_tables,
    sync_documents_from_uploads,
)
from services.solvency_pdf_locator import PageMatch  # noqa: E402
from services.solvency_step3_standardizer import batch_result_workbook_bytes  # noqa: E402
from services.solvency_normalizer import STANDARD_COLUMNS  # noqa: E402
from services.solvency_rule_table_extractor import MODE_RULE, MODE_VISION  # noqa: E402
import services.batch_workflow as batch_workflow  # noqa: E402
from ui.cache import read_profile_registry, profile_registry_token  # noqa: E402
from ui.batch_workspace import batch_status_frame, failed_batch_companies  # noqa: E402
from tests.test_new_workflow_behavior import build_three_table_pdf  # noqa: E402


def _match(
    table_id: str,
    table_name: str,
    content_terms: list[str],
) -> PageMatch:
    return PageMatch(
        table_id=table_id,
        table_name=table_name,
        pages=[1],
        score=1.0,
        evidence="synthetic",
        review_required=False,
        review_reason="",
        sources=[],
        strategy_id=None,
        table_config={"content_terms": content_terms},
        confidence=100,
        confidence_basis="synthetic",
    )


def test_upload_sync_keeps_same_pdf_results() -> None:
    pdf = build_three_table_pdf()
    reads = {"count": 0}

    def _getvalue() -> bytes:
        reads["count"] += 1
        return pdf

    uploads = [
        SimpleNamespace(
            name="甲公司.pdf",
            size=len(pdf),
            file_id="file-1",
            getvalue=_getvalue,
        )
    ]
    documents = sync_documents_from_uploads({}, uploads, profile_id="NON_LIFE_SOLVENCY")
    documents["甲公司"]["locate_status"] = "成功"
    again = sync_documents_from_uploads(
        documents,
        uploads,
        profile_id="NON_LIFE_SOLVENCY",
    )
    assert again["甲公司"]["locate_status"] == "成功"
    assert reads["count"] == 0
    assert document_pdf_bytes(again["甲公司"]) == pdf
    assert reads["count"] == 1
    print("PASS 多文件上传重复 rerun 不丢失已有定位结果")


def test_89_pdf_upload_sync_is_lazy() -> None:
    def _must_not_read() -> bytes:
        raise AssertionError("上传同步阶段不应读取 PDF 正文")

    uploads = [
        SimpleNamespace(
            name=f"公司{index:02d}.pdf",
            size=1024,
            file_id=f"file-{index}",
            getvalue=_must_not_read,
        )
        for index in range(89)
    ]
    documents = sync_documents_from_uploads(
        {},
        uploads,
        profile_id="NON_LIFE_SOLVENCY",
    )
    assert len(documents) == 89
    documents["公司00"]["locate_method"] = "规则法"
    documents["公司00"]["extraction_method"] = "全部规则法（离线，零模型调用）"
    status = batch_status_frame(documents)
    assert len(status) == 89
    assert {"定位逻辑", "提取逻辑"}.issubset(status.columns)
    again = sync_documents_from_uploads(
        documents,
        uploads,
        profile_id="NON_LIFE_SOLVENCY",
    )
    assert len(again) == 89
    print("PASS 89 份 PDF 上传同步零正文读取，状态表显示定位/提取逻辑")


def test_locate_success_without_pages_is_reported_as_failed() -> None:
    documents = {
        "甲公司": {
            "locate_status": "成功",
            "matches": [
                _match("SOLVENCY_MAIN", "偿付能力充足率指标", ["认可资产"])
            ],
        },
        "乙公司": {
            "locate_status": "成功",
            "matches": [
                SimpleNamespace(table_id="SOLVENCY_MAIN", pages=[])
            ],
        },
    }
    assert failed_batch_companies(documents, "locate") == ["乙公司"]
    print("PASS 定位成功但无页码的公司进入失败清单")


def test_batch_rule_extraction_and_standardization() -> None:
    pdf = build_three_table_pdf()
    matches = [
        _match(
            "SOLVENCY_MAIN",
            "偿付能力充足率指标",
            ["认可资产", "综合偿付能力充足率"],
        ),
        _match(
            "LIQUIDITY_MONITORING",
            "流动性风险监测指标",
            ["LCR1"],
        ),
    ]
    extraction = extract_document_tables(
        company="甲公司",
        pdf_bytes=pdf,
        matches=matches,
        method="全部规则法（离线，零模型调用）",
        allow_llm_fallback=False,
        whole_page=True,
    )
    assert extraction.status == "成功", extraction.note
    assert {table.table_id for table in extraction.tables} == {
        "SOLVENCY_MAIN",
        "LIQUIDITY_MONITORING",
    }
    assert all(log["识别方式"] == "规则法" for log in extraction.logs)

    profile = read_profile_registry(profile_registry_token())["NON_LIFE_SOLVENCY"]
    context = build_context_from_profile(
        profile,
        api_key="",
        base_url="",
        model="",
        use_vision=False,
    )
    standardized = standardize_document_tables(
        company="甲公司",
        pdf_name="甲公司.pdf",
        pdf_bytes=pdf,
        tables=extraction.tables,
        context=context,
    )
    assert standardized["状态"] == "成功", standardized["说明"]
    data = standardized["data"]
    assert data["指标编码"].eq("RECOGNIZED_ASSETS").any()
    assert data["指标编码"].eq("COMBINED_SOLVENCY_RATIO").any()
    liquidity = next(
        table for table in extraction.tables
        if table.table_id == "LIQUIDITY_MONITORING"
    )
    assert "LCR1" in "".join(
        str(cell) for row in liquidity.rows for cell in row
    )
    print("PASS 批量默认规则法提取并标准化多个目标表")


def test_merged_workbook_has_exactly_one_sheet() -> None:
    frame = pd.DataFrame(
        [
            {"公司": "甲公司", "指标编码": "A", "数值": 1},
            {"公司": "乙公司", "指标编码": "B", "数值": 2},
        ]
    )
    merged = merge_standard_frames([frame, frame.assign(公司="丙公司")])
    assert merged["公司"].nunique() == 3
    workbook = single_sheet_workbook_bytes(merged, sheet_name="合并指标")
    excel = pd.ExcelFile(io.BytesIO(workbook))
    assert excel.sheet_names == ["合并指标"]
    assert len(pd.read_excel(io.BytesIO(workbook))) == 4
    print("PASS 全部公司指标可下载为单个 Excel 工作表")


def test_scanned_page_routes_to_vision_model_in_batch() -> None:
    import fitz

    source = fitz.open()
    source_page = source.new_page()
    source_page.insert_text((72, 72), "扫描页测试", fontname="china-s")
    pixmap = source_page.get_pixmap(alpha=False)
    image_pdf = fitz.open()
    image_page = image_pdf.new_page(width=source_page.rect.width, height=source_page.rect.height)
    image_page.insert_image(image_page.rect, pixmap=pixmap)
    scanned_bytes = image_pdf.tobytes()
    image_pdf.close()
    source.close()
    match = _match("OPERATING_METRICS", "主要经营指标（财险）", ["保险业务收入"])
    calls: list[str] = []

    def fake_extract(**kwargs):
        calls.append(kwargs["method"])
        return SimpleNamespace(
            frame=None,
            frames=(),
            engine="llm_vision",
            note="mock vision call",
            succeeded=False,
        )

    with patch.object(batch_workflow, "extract_page_with_method", fake_extract):
        batch_workflow.extract_document_tables(
            company="扫描版财险",
            pdf_bytes=scanned_bytes,
            matches=[match],
            method=MODE_RULE,
            vision_for_scanned_pages=True,
            api_key="test-key",
            base_url="https://example.invalid/v1",
            model="vision-test",
        )
    assert calls and calls[0] == MODE_VISION
    print("PASS 批量图片型页面自动转视觉大模型")


def test_batch_workbook_uses_the_single_flow_format() -> None:
    from openpyxl import load_workbook

    data = pd.DataFrame(columns=STANDARD_COLUMNS)
    summary = pd.DataFrame(
        [
            {
                "公司": "甲公司",
                "指标编码": "SIGNED_PREMIUM",
                "指标名称": "签单保费",
                "填报状态": "已填报",
            },
            {
                "公司": "甲公司",
                "指标编码": "NET_PROFIT",
                "指标名称": "净利润",
                "填报状态": "未填报",
            },
        ]
    )
    payload = batch_result_workbook_bytes(
        data,
        summary,
        pd.DataFrame([{"状态": "未匹配", "原因": "测试"}]),
    )
    workbook = load_workbook(io.BytesIO(payload), read_only=True)
    assert "标准数据" in workbook.sheetnames
    assert "填报汇总" in workbook.sheetnames
    assert "匹配诊断" in workbook.sheetnames


def test_batch_workspace_renders_step1_to_step4() -> None:
    pdf = build_three_table_pdf()
    matches = [
        _match(
            "SOLVENCY_MAIN",
            "偿付能力充足率指标",
            ["认可资产", "综合偿付能力充足率"],
        )
    ]
    extraction = extract_document_tables(
        company="甲公司",
        pdf_bytes=pdf,
        matches=matches,
    )
    profile = read_profile_registry(profile_registry_token())["NON_LIFE_SOLVENCY"]
    context = build_context_from_profile(
        profile,
        api_key="",
        base_url="",
        model="",
        use_vision=False,
    )
    standardized = standardize_document_tables(
        company="甲公司",
        pdf_name="甲公司.pdf",
        pdf_bytes=pdf,
        tables=extraction.tables,
        context=context,
    )
    document = {
        "company": "甲公司",
        "pdf_name": "甲公司.pdf",
        "pdf_bytes": pdf,
        "profile_id": "NON_LIFE_SOLVENCY",
        "metadata": {"页数": 1},
        "matches": matches,
        "edited_pages": {},
        "pages_confirmed": True,
        "locate_status": "成功",
        "locate_note": "",
        "raw_tables": extraction.tables,
        "extraction_logs": extraction.logs,
        "extraction_status": extraction.status,
        "extraction_note": extraction.note,
        "standard_data": standardized["data"],
        "standard_summary": standardized["summary"],
        "standard_diagnostics": standardized["diagnostics"],
        "standard_workbook_bytes": standardized["workbook_bytes"],
        "standard_status": standardized["状态"],
        "standard_note": standardized["说明"],
    }
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=240)
    at.session_state["logged_in"] = True
    at.session_state["user_role"] = "项目组成员"
    at.session_state["active_profile_id"] = "NON_LIFE_SOLVENCY"
    at.session_state["batch_documents"] = {"甲公司": document}
    at.session_state["standard_data"] = standardized["data"]
    for label in (
        "STEP1 页码定位",
        "STEP2 表格提取",
        "STEP3 标准化",
        "STEP4 勾稽检查",
    ):
        at.session_state["workflow_main_tab"] = label
        at.run()
        assert not at.exception, (label, [item.value for item in at.exception])
        if label == "STEP1 页码定位":
            mode_radios = [
                item for item in at.radio
                if item.key and item.key.startswith("batch_locate_mode_")
            ]
            assert mode_radios
            assert any("混合法" in option for option in mode_radios[0].options)
            assert any(
                item.key == "batch_relocate_companies_NON_LIFE_SOLVENCY"
                for item in at.multiselect
            )
        if label == "STEP2 表格提取":
            assert any(
                item.key == "batch_retry_companies_NON_LIFE_SOLVENCY"
                for item in at.multiselect
            )
    print("PASS 批量文件池在 STEP1-STEP4 页面正常渲染")


if __name__ == "__main__":
    test_upload_sync_keeps_same_pdf_results()
    test_89_pdf_upload_sync_is_lazy()
    test_locate_success_without_pages_is_reported_as_failed()
    test_batch_rule_extraction_and_standardization()
    test_merged_workbook_has_exactly_one_sheet()
    test_batch_workbook_uses_the_single_flow_format()
    test_scanned_page_routes_to_vision_model_in_batch()
    test_batch_workspace_renders_step1_to_step4()
    print("全部通过")
