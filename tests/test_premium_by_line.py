# -*- coding: utf-8 -*-
"""分险种保费收入：集成表读取、同比增速、分行柱状/折线图与中英文渲染回归。"""

from __future__ import annotations

import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from openpyxl import load_workbook  # noqa: E402

from services.solvency_step7_charts import (  # noqa: E402
    DISPLAY_LINES,
    INTEGRATED_SHEET_NAME,
    INTEGRATED_WORKBOOK_PATH,
    PREMIUM_BY_LINE_PATH,
    build_premium_by_line_combo_chart,
    load_premium_by_line_data,
    prepare_premium_by_line_overview,
    premium_by_line_chart_frame,
    premium_by_line_display_frame,
    premium_by_line_source_key,
    premium_by_line_window_years,
    premium_line_frame,
)
from services.report_i18n import localize_chart  # noqa: E402
from services.step7_excel_config import load_step7_config  # noqa: E402


def test_integrated_workbook_contains_sheet() -> None:
    assert INTEGRATED_WORKBOOK_PATH.exists(), "未找到行业集成数据工作簿"
    workbook = load_workbook(
        INTEGRATED_WORKBOOK_PATH,
        read_only=True,
        data_only=True,
    )
    try:
        assert INTEGRATED_SHEET_NAME in workbook.sheetnames
        sheet = workbook[INTEGRATED_SHEET_NAME]
        assert sheet.max_row == 52  # 表头 + 51 条记录
        assert sheet.max_column == 7
    finally:
        workbook.close()
    assert premium_by_line_source_key() == "integrated"


def test_source_data_is_loaded() -> None:
    assert PREMIUM_BY_LINE_PATH.exists()
    frame = load_premium_by_line_data()
    assert len(frame) == 51
    assert set(frame.columns) == {"年度", "险种", "原保险保费收入", "同比增速", "基期年度"}
    assert frame["年度"].min() == 2019
    assert frame["年度"].max() == 2025
    assert frame["险种"].nunique() == 9


def test_growth_uses_adjacent_year_with_data() -> None:
    frame = load_premium_by_line_data()
    motor = premium_line_frame(frame, "机动车辆保险")
    assert motor["年度"].tolist() == list(range(2019, 2026))
    latest = motor.iloc[-1]
    assert latest["原保险保费收入"] == 9409.0
    assert round(float(latest["同比增速"]), 2) == round((9409 / 9137 - 1) * 100, 2)
    assert latest["基期年度"] == "2024"

    engineering = premium_line_frame(frame, "工程保险")
    assert engineering["年度"].tolist() == [2019, 2020, 2021, 2022]
    assert engineering.iloc[-1]["同比增速"] is not None


def test_display_frame_keeps_five_lines_and_latest_four_years() -> None:
    frame = load_premium_by_line_data()
    years = premium_by_line_window_years(frame)
    assert years == [2022, 2023, 2024, 2025]
    display = premium_by_line_display_frame(frame, years=years)
    assert len(display) == 5 * 4
    assert set(display["险种"].astype(str)) == set(DISPLAY_LINES)
    assert sorted(display["年度"].unique().tolist()) == years
    # 过滤后第一年仍保留按完整历史计算的同比与基期。
    motor_2022 = display[
        display["险种"].eq("机动车辆保险") & display["年度"].eq(2022)
    ].iloc[0]
    assert motor_2022["基期年度"] == "2021"
    assert round(float(motor_2022["同比增速"]), 2) == round((8210 / 7773 - 1) * 100, 2)


def test_chart_frame_marks_missing_years_for_retired_lines() -> None:
    frame = load_premium_by_line_data()
    display = premium_by_line_display_frame(frame)
    chart_frame = premium_by_line_chart_frame(display)
    assert len(chart_frame) == 5 * 4
    assert chart_frame["披露状态"].eq("未披露").sum() == 0

    engineering = premium_by_line_chart_frame(frame, line_name="工程保险")
    assert len(engineering) == 7  # 2019–2025
    missing = engineering[engineering["披露状态"].eq("未披露")]
    assert set(missing["年度"]) == {"2023", "2024", "2025"}


def test_combo_chart_layers_premium_bars_and_growth_line() -> None:
    frame = load_premium_by_line_data()
    years = premium_by_line_window_years(frame)
    chart = build_premium_by_line_combo_chart(
        premium_by_line_display_frame(frame, years=years),
        "机动车辆保险",
        years=years,
    )
    spec = chart.to_dict()
    assert spec["width"] == "container"
    assert spec["resolve"]["scale"]["y"] == "independent"
    assert [layer["mark"]["type"] for layer in spec["layer"]] == [
        "bar",
        "text",
        "rule",
        "line",
        "text",
    ]
    bars = spec["layer"][0]
    assert bars["encoding"]["x"]["field"] == "年度"
    assert "xOffset" not in bars["encoding"]
    assert bars["encoding"]["y"]["field"] == "原保险保费收入"
    assert bars["encoding"]["y"]["title"] == "原保险保费收入（亿元）"
    line = spec["layer"][3]
    assert line["encoding"]["y"]["field"] == "同比增速"
    assert line["encoding"]["y"]["title"] == "同比增速（%）"
    assert line["encoding"]["y"]["axis"]["labelExpr"] == "datum.label + '%'"
    assert line["encoding"]["y"]["axis"]["format"] == ".0f"
    assert line["encoding"]["y"]["axis"]["orient"] == "right"
    assert line["mark"]["point"]["filled"] is True
    assert bars["encoding"]["y"]["axis"]["orient"] == "left"
    # 辅助文本层不能再生成第二条金额/增速轴，否则右轴会出现两条刻度重叠。
    y_axes = [layer["encoding"]["y"].get("axis") for layer in spec["layer"]]
    assert y_axes[1] is None and y_axes[2] is None and y_axes[4] is None
    assert sum(axis is not None for axis in y_axes) == 2
    # 折线与柱体使用不同颜色，避免同色导致走势看不清。
    assert line["mark"]["color"] != bars["mark"]["color"]
    assert line["mark"]["color"] == "#C2410C"


def test_english_localisation_has_no_chinese() -> None:
    frame = load_premium_by_line_data()
    years = premium_by_line_window_years(frame)
    display = premium_by_line_display_frame(frame, years=years)
    chart = build_premium_by_line_combo_chart(
        display, "机动车辆保险", years=years
    )
    spec = localize_chart(chart, "en").to_dict()
    text = str(spec)
    assert "Motor vehicle insurance" in text
    rows: list[dict] = []
    for dataset_rows in spec.get("datasets", {}).values():
        if isinstance(dataset_rows, list):
            rows.extend(row for row in dataset_rows if isinstance(row, dict))
    assert any(row.get("险种") == "Motor vehicle insurance" for row in rows)
    assert "Original premium income (RMB 100 mn)" in text
    assert "YoY growth (%)" in text


def test_overview_uses_filtered_lines() -> None:
    frame = load_premium_by_line_data()
    display = premium_by_line_display_frame(frame)
    overview = prepare_premium_by_line_overview(display)
    assert overview is not None
    assert overview.year == 2025
    assert overview.total == 15259.0
    assert overview.previous_total == 14570.0
    assert round(overview.growth, 2) == round((15259 / 14570 - 1) * 100, 2)
    assert overview.top_line == "机动车辆保险"
    assert overview.fastest_line == "健康险"
    assert overview.legacy_lines == ()


def test_template_places_section_between_scale_and_market_share() -> None:
    modules = load_step7_config(
        ROOT / "templates" / "Step7-偿付能力公司报告展示模块.xlsx"
    )
    industry = [module for module in modules if module.level_one == "行业概览"]
    assert [module.level_two for module in industry] == [
        "规模与发展",
        "分险种保费收入",
        "市场份额",
        "承保效益",
        "偿付能力披露整体情况",
    ]
    by_line = industry[1]
    assert by_line.chart_name == "分险种保费收入"
    assert by_line.chart_kind == "industry_premium_by_line"


def _app(level_two: str, language: str):
    from streamlit.testing.v1 import AppTest

    from test_workflow_smoke import synthetic_integrated_data

    template_path = ROOT / "templates" / "Step7-偿付能力公司报告展示模块.xlsx"
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=900)
    at.session_state["logged_in"] = True
    at.session_state["user_role"] = "项目组成员"
    at.session_state["active_profile_id"] = "NON_LIFE_SOLVENCY"
    data = synthetic_integrated_data()
    at.session_state["integrated_data"] = data
    at.session_state["standard_data"] = data
    at.session_state["step7_notes_bytes_NON_LIFE_SOLVENCY"] = template_path.read_bytes()
    at.session_state["workflow_main_tab"] = "STEP7 公司报告"
    at.session_state["company_nav_level_one"] = "行业概览"
    at.session_state["company_nav_level_two"] = level_two
    at.session_state["s7_company_types"] = ["全部"]
    at.session_state["s7_unit_mode"] = "亿元"
    at.session_state["s7_language"] = language
    at.session_state["s7_language_display"] = "English" if language == "en" else "中文"
    return at


def test_chinese_page_renders_five_combo_charts() -> None:
    at = _app("分险种保费收入", "zh")
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    markdown_text = "\n".join(item.value for item in at.markdown)
    caption_text = "\n".join(item.value for item in at.caption)
    combined = markdown_text + "\n" + caption_text
    assert "分险种原保险保费收入与同比增速" in markdown_text
    assert "每行两个险种" in markdown_text
    for name in ("机动车辆保险", "健康险", "农业保险", "责任险", "意外险"):
        assert name in markdown_text
    assert "原保险保费收入（左轴，亿元）" in combined
    assert "同比增速（右轴，%）" in combined
    assert len(at.get("vega_lite_chart")) == 5
    print("PASS 分险种保费收入：中文页面五个险种各一张组合图正常")


def test_english_page_has_no_chinese_leak() -> None:
    at = _app("分险种保费收入", "en")
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    markdown_text = "\n".join(item.value for item in at.markdown)
    caption_text = "\n".join(item.value for item in at.caption)
    combined = markdown_text + "\n" + caption_text
    assert "Original premium income and YoY growth by line" in combined
    for name in (
        "Motor vehicle insurance",
        "Health insurance",
        "Agricultural insurance",
        "Liability insurance",
        "Accident insurance",
    ):
        assert name in combined
    assert len(at.get("vega_lite_chart")) == 5

    chart_text = ""
    for chart in at.get("vega_lite_chart"):
        chart_text += str(chart.proto.spec)
        chart_text += "\n".join(str(dataset.data) for dataset in chart.proto.datasets)

    combined = re.sub(r"<style[\s\S]*?</style>", "", combined)
    combined = re.sub(
        r'<section class="platform-page-heading"[\s\S]*?</section>',
        "",
        combined,
    )
    leaked = [
        line.strip()
        for line in (combined + "\n" + chart_text).splitlines()
        if re.search(r"[\u4e00-\u9fff]", line)
        and not any(
            marker in line
            for marker in (
                "当前身份：",
                "Profile：",
                "默认本地 Profile：",
                "目标表和定位关键词",
                "### 行业分析导航",
                "切换后会清空",
                "v2 配置：",
                "配置文件保存",
                "当前配置范围：",
            )
        )
    ]
    assert not leaked, "\n".join(leaked[:30])
    print("PASS 分险种保费收入：英文页面完成切换且无中文残留")


if __name__ == "__main__":
    test_integrated_workbook_contains_sheet()
    test_source_data_is_loaded()
    test_growth_uses_adjacent_year_with_data()
    test_display_frame_keeps_five_lines_and_latest_four_years()
    test_chart_frame_marks_missing_years_for_retired_lines()
    test_combo_chart_layers_premium_bars_and_growth_line()
    test_english_localisation_has_no_chinese()
    test_overview_uses_filtered_lines()
    test_template_places_section_between_scale_and_market_share()
    test_chinese_page_renders_five_combo_charts()
    test_english_page_has_no_chinese_leak()
    print("全部通过")
