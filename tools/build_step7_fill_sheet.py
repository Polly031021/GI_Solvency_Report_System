"""Generate a paste-ready STEP7 missing-metric fill sheet.

The output workbook contains a sheet with exactly the same columns/order as the
source sheet (default: 标准数据). It contains only missing metric rows, with the
value column left blank. Fill the values, then copy the data rows and paste them
at the bottom of the existing table.

Usage:
    python tools/build_step7_fill_sheet.py <source.xlsx>
    python tools/build_step7_fill_sheet.py <source.xlsx> --scope core
    python tools/build_step7_fill_sheet.py <source.xlsx> --all-periods
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from services.solvency_step7_metric_catalog import (  # noqa: E402
    step7_metric_requirement_catalog,
)
from services.step3_metric_gap import build_step7_metric_gap  # noqa: E402


CORE_CODES = {
    "ACTUAL_CAPITAL",
    "MINIMUM_CAPITAL",
    "CORE_T1_CAPITAL",
    "CORE_T2_CAPITAL",
    "ANC_T1_CAPITAL",
    "ANC_T2_CAPITAL",
    "COMBINED_SOLVENCY_RATIO",
    "CORE_SOLVENCY_RATIO",
    "SIGNED_PREMIUM",
    "SIGNED_PREMIUM_MOTOR",
    "NET_PROFIT",
    "COMBINED_RATIO",
    "TOTAL_ASSETS",
    "RECOGNIZED_ASSETS",
    "RECOGNIZED_LIABILITIES",
    "REGISTERED_CAPITAL",
}

OUTPUT_SHEET = "标准数据"
SUMMARY_SHEET = "缺失指标汇总"
HELP_SHEET = "填写说明"
CATALOG_SHEET = "STEP7指标目录"

_PERIOD_RE = re.compile(r"(20\d{2})\D*[Qq]?([1-4])")
_COMPANY_SUFFIX_RE = re.compile(
    r"[（(]\s*(?:偿付能力|财险|寿险|健康险|养老险|合并|母公司|集团)\s*[)）]\s*$"
)


def _text(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).strip()


def _company_key(value: object) -> str:
    return _COMPANY_SUFFIX_RE.sub("", _text(value)).strip()


def _period_sort(value: object) -> tuple[int, int, str]:
    text = _text(value)
    match = _PERIOD_RE.search(text)
    if match:
        return int(match.group(1)), int(match.group(2)), text
    return 0, 0, text


def _report_parts(value: object) -> tuple[str, str]:
    match = _PERIOD_RE.search(_text(value))
    if not match:
        return "", ""
    return match.group(1), f"Q{match.group(2)}"


def _read_source(path: Path, sheet_name: str) -> pd.DataFrame:
    frame = pd.read_excel(path, sheet_name=sheet_name, dtype=str)
    frame.columns = [str(column).strip() for column in frame.columns]
    for column in ("公司", "报告期", "指标编码", "指标名称", "期间口径", "单位", "数据类型"):
        if column not in frame.columns:
            raise ValueError(f"{sheet_name} 缺少必要列：{column}")
        frame[column] = frame[column].fillna("").astype(str).str.strip()
    frame["_公司键"] = frame["公司"].map(_company_key)
    frame["_期序"] = frame["报告期"].map(_period_sort)
    return frame


def _company_code_map(path: Path) -> dict[str, str]:
    try:
        internal = pd.read_excel(
            path,
            sheet_name="_系统字段",
            dtype=str,
            usecols=["公司", "公司统一编码"],
        )
    except Exception:
        return {}
    internal.columns = [str(column).strip() for column in internal.columns]
    internal["公司"] = internal["公司"].fillna("").astype(str).str.strip()
    internal["公司统一编码"] = internal["公司统一编码"].fillna("").astype(str).str.strip()
    internal = internal[(internal["公司"].ne("")) & (internal["公司统一编码"].ne(""))]
    return internal.drop_duplicates("公司", keep="first").set_index("公司")["公司统一编码"].to_dict()


def _latest_subset(frame: pd.DataFrame) -> pd.DataFrame:
    latest = (
        frame.sort_values(["公司", "_期序"], kind="stable")
        .groupby("公司", as_index=False)
        .tail(1)[["公司", "报告期"]]
        .drop_duplicates(["公司", "报告期"])
    )
    return frame.merge(latest, on=["公司", "报告期"], how="inner")


def _meta_lookup(catalog: pd.DataFrame) -> dict[str, dict]:
    lookup: dict[str, dict] = {}
    for _, row in catalog.iterrows():
        code = _text(row.get("指标编码"))
        if code:
            lookup[code] = row.to_dict()
    return lookup


def _missing_frame(
    data: pd.DataFrame,
    *,
    scope: str,
    catalog: pd.DataFrame,
) -> tuple[pd.DataFrame, object]:
    report = build_step7_metric_gap(data, requirement_catalog=catalog)
    missing = report.missing.copy()
    if scope == "core":
        missing = missing[missing["指标编码"].isin(CORE_CODES)].copy()
    return missing, report


def _paste_rows(
    missing: pd.DataFrame,
    source: pd.DataFrame,
    columns: list[str],
    catalog: pd.DataFrame,
    code_map: dict[str, str],
) -> pd.DataFrame:
    meta = _meta_lookup(catalog)
    source_lookup: dict[tuple[str, str], pd.Series] = {}
    for _, row in source.drop_duplicates(["_公司键", "报告期"], keep="first").iterrows():
        source_lookup[(_text(row["_公司键"]), _text(row["报告期"]))] = row

    records: list[dict] = []
    for _, miss in missing.iterrows():
        company_key = _company_key(miss.get("公司"))
        report_period = _text(miss.get("报告期"))
        base = source_lookup.get((company_key, report_period))
        if base is None:
            base_rows = source[source["_公司键"].eq(company_key)]
            base = base_rows.iloc[0] if not base_rows.empty else pd.Series(dtype=object)

        record = {column: base.get(column, "") for column in columns}
        code = _text(miss.get("指标编码"))
        info = meta.get(code, {})
        category = "核心必备" if code in CORE_CODES else "明细/条件"
        year, quarter = _report_parts(report_period)

        record.update(
            {
                "公司": _text(miss.get("公司")) or _text(base.get("公司")),
                "报告期": report_period,
                "一级模块": _text(info.get("一级模块")),
                "二级模块": _text(info.get("二级模块")),
                "指标编码": code,
                "指标名称": _text(info.get("指标名称")) or _text(miss.get("指标名称")),
                "期间口径": _text(miss.get("期间口径")) or _text(info.get("建议期间口径")),
                "数值": None,
                "单位": _text(info.get("标准单位")) or _text(miss.get("单位")),
                "数据类型": _text(info.get("数据类型")) or _text(miss.get("数据类型")),
                "指标属性": _text(info.get("指标属性")) or "披露",
                "计算逻辑": _text(info.get("计算逻辑")),
                "备注": f"STEP7缺失补录（{category}）；填写数值后粘贴到主表末尾",
                "来源类型": "人工补录",
                "来源工作表": "STEP7缺失补录",
                "来源页码": "",
                "原始披露值": "",
                "是否预测": "否",
                "行次": "",
                "报告年度": _text(base.get("报告年度")) or year,
                "报告季度": _text(base.get("报告季度")) or quarter,
            }
        )
        if "公司统一编码" in columns and not _text(record.get("公司统一编码")):
            record["公司统一编码"] = code_map.get(_text(record.get("公司")), "")
        records.append(record)

    output = pd.DataFrame(records, columns=columns)
    if not output.empty:
        output["_类别序"] = output["指标编码"].map(
            lambda code: 0 if code in CORE_CODES else 1
        )
        output = (
            output.sort_values(["公司", "_类别序", "指标编码"], kind="stable")
            .drop(columns=["_类别序"])
            .reset_index(drop=True)
        )
    return output


def _summary_frame(output: pd.DataFrame) -> pd.DataFrame:
    if output.empty:
        return pd.DataFrame(columns=["指标编码", "指标名称", "指标类别", "缺失公司数", "缺失记录数"])
    summary = (
        output.groupby(["指标编码", "指标名称"], dropna=False)
        .agg(
            缺失记录数=("公司", "size"),
            缺失公司数=("公司", "nunique"),
        )
        .reset_index()
    )
    summary["指标类别"] = summary["指标编码"].map(
        lambda code: "核心必备" if code in CORE_CODES else "明细/条件"
    )
    summary = summary[["指标编码", "指标名称", "指标类别", "缺失公司数", "缺失记录数"]]
    return summary.sort_values(
        ["指标类别", "缺失公司数", "指标编码"],
        ascending=[True, False, True],
    )


def _help_frame(
    source: Path,
    output: pd.DataFrame,
    sheet_name: str,
    scope: str,
    periods: str,
    total_companies: int,
) -> pd.DataFrame:
    core_count = int(output["指标编码"].isin(CORE_CODES).sum()) if not output.empty else 0
    return pd.DataFrame(
        [
            ("源文件", source.name),
            ("源工作表", sheet_name),
            ("补录范围", "核心必备指标" if scope == "core" else "全部 STEP7 基础指标"),
            ("报告期范围", periods),
            ("数据行数", f"{len(output):,} 行（仅缺失指标，不含已有数据）"),
            ("涉及公司", f"{output['公司'].nunique() if not output.empty else 0} 家 / 源表 {total_companies} 家"),
            ("核心缺口行数", f"{core_count:,} 行"),
            (
                "使用步骤",
                "1. 在“标准数据”页的黄色“数值”列填写数字；"
                "2. 选中数据区（不要选表头），复制；"
                "3. 到已有汇总表的最后一行下方，右键“粘贴为值”；"
                "4. 粘贴后可用筛选检查“指标编码”是否已补齐。",
            ),
            (
                "单位",
                "金额标准单位为万元；若原始数据是元，把“单位”改为“元”，"
                "粘贴后由后续标准化流程换算；不要同时改数值和单位。",
            ),
            (
                "期间口径",
                "默认按指标类型预填：存量指标本季度末数，流量指标本年累计数；"
                "如实际披露口径不同，可直接在“期间口径”列修改。",
            ),
            (
                "注意",
                "本表列顺序与源工作表完全一致，可直接粘贴；如某行不适用，删除该行后再粘贴。",
            ),
        ],
        columns=["项目", "说明"],
    )


def _style_sheet(sheet, *, value_column: int | None, unit_column: int | None, period_column: int | None) -> None:
    header_fill = PatternFill("solid", fgColor="00338D")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10)
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sheet.freeze_panes = "B2"
    if sheet.max_row >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    sheet.row_dimensions[1].height = 34
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    for column_cells in sheet.iter_cols(min_row=1, max_row=min(sheet.max_row, 400)):
        values = []
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            values.append(sum(2 if ord(char) > 127 else 1 for char in value))
        width = min(max(max(values, default=8) + 3, 10), 42)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

    if value_column is not None:
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=value_column)
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
            cell.number_format = "#,##0.00####"
            cell.alignment = Alignment(horizontal="right", vertical="top")

        sheet.cell(row=1, column=value_column).comment = Comment(
            "在这里填写数字；金额按“单位”列口径，标准单位为万元。",
            "STEP7",
        )
        value_letter = get_column_letter(value_column)
        validation = DataValidation(
            type="decimal",
            operator="between",
            formula1="-1000000000000000",
            formula2="1000000000000000",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="数值格式",
            error="请填写数字，不要带千分位或单位文字。",
        )
        sheet.add_data_validation(validation)
        validation.add(f"{value_letter}2:{value_letter}{max(sheet.max_row, 2)}")

    if unit_column:
        unit_letter = get_column_letter(unit_column)
        unit_validation = DataValidation(
            type="list",
            formula1='"元,千元,万元,百元,亿元,十亿元,%,倍"',
            allow_blank=True,
        )
        sheet.add_data_validation(unit_validation)
        unit_validation.add(f"{unit_letter}2:{unit_letter}{max(sheet.max_row, 2)}")

    if period_column:
        period_letter = get_column_letter(period_column)
        period_validation = DataValidation(
            type="list",
            formula1='"本季度末数,上季度末数,本季度数,上季度数,本年累计数,近三年平均,期末数,期初数"',
            allow_blank=True,
        )
        sheet.add_data_validation(period_validation)
        period_validation.add(f"{period_letter}2:{period_letter}{max(sheet.max_row, 2)}")


def build_fill_sheet(
    source: Path,
    output: Path,
    *,
    sheet_name: str = "标准数据",
    scope: str = "all",
    all_periods: bool = False,
) -> Path:
    source_frame = _read_source(source, sheet_name)
    work = source_frame[source_frame["指标编码"].ne("")].copy()
    data = work if all_periods else _latest_subset(work)
    catalog = step7_metric_requirement_catalog()
    missing, report = _missing_frame(data, scope=scope, catalog=catalog)
    columns = list(source_frame.columns)
    # 内部辅助列不进入粘贴表
    columns = [column for column in columns if not column.startswith("_")]
    code_map = _company_code_map(source)
    output_frame = _paste_rows(missing, data, columns, catalog, code_map)
    summary = _summary_frame(output_frame)
    periods = (
        "全部已有报告期"
        if all_periods
        else "、".join(
            sorted(
                {
                    _text(value)
                    for value in data["报告期"].dropna().astype(str).unique()
                }
            )
        )
    )
    help_frame = _help_frame(
        source,
        output_frame,
        sheet_name,
        scope,
        periods,
        int(work["公司"].nunique()),
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        output_frame.to_excel(writer, sheet_name=OUTPUT_SHEET, index=False)
        summary.to_excel(writer, sheet_name=SUMMARY_SHEET, index=False)
        help_frame.to_excel(writer, sheet_name=HELP_SHEET, index=False)
        catalog.to_excel(writer, sheet_name=CATALOG_SHEET, index=False)

        paste_sheet = writer.book[OUTPUT_SHEET]
        value_column = columns.index("数值") + 1
        unit_column = columns.index("单位") + 1 if "单位" in columns else None
        period_column = (
            columns.index("期间口径") + 1 if "期间口径" in columns else None
        )
        _style_sheet(
            paste_sheet,
            value_column=value_column,
            unit_column=unit_column,
            period_column=period_column,
        )
        for name in (SUMMARY_SHEET, HELP_SHEET, CATALOG_SHEET):
            sheet = writer.book[name]
            _style_sheet(
                sheet,
                value_column=None,
                unit_column=None,
                period_column=None,
            )

    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="含标准数据页的汇总工作簿")
    parser.add_argument("output", type=Path, nargs="?", help="输出补录工作簿")
    parser.add_argument("--sheet", default="标准数据", help="源工作表名，默认“标准数据”")
    parser.add_argument(
        "--scope",
        choices=("all", "core"),
        default="all",
        help="all=全部 STEP7 基础指标；core=仅核心必备指标",
    )
    parser.add_argument(
        "--all-periods",
        action="store_true",
        help="检查全部已有报告期；默认每家公司只取最后一个报告期",
    )
    args = parser.parse_args()

    output = args.output
    if output is None:
        suffix = "_核心" if args.scope == "core" else "_全部"
        output = args.source.with_name(
            f"{args.source.stem}_STEP7缺失指标_可粘贴补录表{suffix}.xlsx"
        )

    result = build_fill_sheet(
        args.source,
        output,
        sheet_name=args.sheet,
        scope=args.scope,
        all_periods=args.all_periods,
    )
    print(f"Wrote {result}")


if __name__ == "__main__":
    main()
