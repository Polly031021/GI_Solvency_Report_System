from __future__ import annotations

"""STEP3 标准化后的 STEP7 指标缺口检查、补录表导出与回填合并。

设计目标：
1. 以 STEP7 指标需求目录为准，检查每家公司/报告期缺少哪些基础指标；
2. 导出标准化的补录 Excel，数值列留空，可直接填写；
3. 回填后按 公司统一编码 + 报告期 + 指标编码 + 期间口径 合并，
   并自动补算派生指标，得到完整的标准窄表。
"""

import io
import re
from dataclasses import dataclass
from typing import Sequence

import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .solvency_dataset_adapter import add_missing_derived_metrics
from .solvency_normalizer import STANDARD_COLUMNS
from .solvency_step3_standardizer import convert_amount_to_target_unit
from .solvency_step7_metric_catalog import (
    DEFAULT_AMOUNT_PERIOD,
    step7_metric_requirement_catalog,
)


MISSING_SHEET_NAME = "缺失指标清单"
DERIVED_GAP_SHEET_NAME = "派生指标缺口"
COVERAGE_SHEET_NAME = "公司覆盖汇总"
INSTRUCTION_SHEET_NAME = "填报说明"
REQUIREMENT_SHEET_NAME = "Step7指标需求"

MISSING_COLUMNS = [
    "公司统一编码",
    "公司",
    "原始公司名称",
    "标准公司名称",
    "公司类型",
    "同业分类",
    "报告类型",
    "报告年度",
    "报告季度",
    "报告期",
    "指标编码",
    "指标名称",
    "一级模块",
    "二级模块",
    "数据类型",
    "单位",
    "期间口径",
    "数值",
    "补录优先级",
    "备注",
]

_COVERAGE_COLUMNS = [
    "公司",
    "公司统一编码",
    "同业分类",
    "报告期",
    "必需基础指标数",
    "已提取基础指标数",
    "缺失基础指标数",
    "覆盖率",
    "派生指标待计算数",
    "缺失指标编码",
]

_DERIVED_GAP_COLUMNS = [
    "公司",
    "公司统一编码",
    "报告期",
    "派生指标编码",
    "派生指标名称",
    "缺失依赖",
    "说明",
]

_COMPANY_SUFFIX_RE = re.compile(
    r"[（(]\s*(?:偿付能力|财险|寿险|健康险|养老险|合并|母公司|集团)\s*[)）]\s*$"
)
_REPORT_PERIOD_RE = re.compile(r"(20\d{2})\s*[-_/]?\s*[Qq]?\s*([1-4])")

_CURRENCY_UNITS = {"元", "千元", "万元", "百元", "亿元", "十亿元"}
_PERIOD_SUGGESTIONS = (
    "本季度末数",
    "上季度末数",
    "本季度数",
    "上季度数",
    "本年累计数",
    "近三年平均",
    "期末数",
    "期初数",
)


def _text(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).strip()


def _company_key(value: object) -> str:
    return _COMPANY_SUFFIX_RE.sub("", _text(value)).strip()


def _to_number(value: object) -> float | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value)
    if not text:
        return None
    text = text.replace("，", ",").replace(",", "").replace(" ", "")
    text = text.replace("％", "%")
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    text = text.replace("%", "")
    try:
        number = float(text)
    except (TypeError, ValueError):
        return None
    return -number if negative else number


def _report_parts(period: object) -> tuple[str, str]:
    match = _REPORT_PERIOD_RE.search(_text(period))
    if not match:
        return "", ""
    return match.group(1), f"Q{match.group(2)}"


def _period_group(value: object) -> str:
    """Normalize period labels for strict coverage checks."""
    text = _text(value)
    base = re.split(r"[-—–]", text)[0].strip()
    if not base:
        return ""
    if "预测" in base:
        return "下季度末预测数"
    if "本年" in base or "累计" in base:
        return "本年累计数"
    if "上季度" in base or "期初" in base or "上年" in base:
        return "上季度末数"
    if "本季度" in base or "期末" in base:
        return "本季度末数"
    return base


@dataclass
class Step7MetricGapReport:
    missing: pd.DataFrame
    coverage: pd.DataFrame
    derived_gaps: pd.DataFrame
    requirement_catalog: pd.DataFrame
    warnings: tuple[str, ...] = ()

    @property
    def missing_count(self) -> int:
        return int(len(self.missing))

    @property
    def company_count(self) -> int:
        return int(self.coverage["公司"].nunique()) if not self.coverage.empty else 0

    @property
    def missing_code_count(self) -> int:
        return (
            int(self.missing["指标编码"].nunique())
            if not self.missing.empty
            else 0
        )

    @property
    def required_base_count(self) -> int:
        if self.requirement_catalog.empty:
            return 0
        return int(
            self.requirement_catalog["补录方式"].eq("直接补录").sum()
        )

    @property
    def coverage_ratio(self) -> float:
        if self.coverage.empty:
            return 1.0
        required = float(self.coverage["必需基础指标数"].sum())
        present = float(self.coverage["已提取基础指标数"].sum())
        return present / required if required else 1.0

    def missing_summary(self, limit: int = 12) -> pd.DataFrame:
        if self.missing.empty:
            return pd.DataFrame(columns=["指标编码", "指标名称", "缺失公司数"])
        summary = (
            self.missing.groupby(["指标编码", "指标名称"], dropna=False)
            .agg(缺失公司数=("公司", "nunique"))
            .reset_index()
            .sort_values(["缺失公司数", "指标编码"], ascending=[False, True])
        )
        return summary.head(limit)


@dataclass
class MissingFillParseResult:
    rows: pd.DataFrame
    skipped: pd.DataFrame
    warnings: tuple[str, ...] = ()

    @property
    def filled_count(self) -> int:
        return int(len(self.rows))


@dataclass
class MissingFillMergeResult:
    data: pd.DataFrame
    added_count: int
    skipped_count: int
    warnings: tuple[str, ...] = ()


def _prepare_data(data: pd.DataFrame | None) -> pd.DataFrame:
    if not isinstance(data, pd.DataFrame) or data.empty:
        return pd.DataFrame(columns=STANDARD_COLUMNS)
    frame = data.copy()
    for column in STANDARD_COLUMNS:
        if column not in frame.columns:
            frame[column] = ""
    frame["公司"] = frame["公司"].fillna("").astype(str).str.strip()
    frame["报告期"] = frame["报告期"].fillna("").astype(str).str.strip()
    frame["指标编码"] = frame["指标编码"].fillna("").astype(str).str.strip()
    frame["期间口径"] = frame["期间口径"].fillna("").astype(str).str.strip()
    frame["_公司键"] = frame["公司"].map(_company_key)
    frame["_数值"] = pd.to_numeric(frame["数值"], errors="coerce")
    return frame


def _pair_records(
    frame: pd.DataFrame,
    company_periods: Sequence[dict] | None,
) -> list[dict]:
    pairs: dict[tuple[str, str], dict] = {}
    if not frame.empty:
        for (company_key, report_period), group in frame.groupby(
            ["_公司键", "报告期"],
            dropna=False,
            sort=False,
        ):
            key = (str(company_key), str(report_period or ""))
            if key[0] in {"行业合计", "合计", "行业总计", "总计"}:
                continue
            first = group.iloc[0]
            pairs[key] = {
                "公司键": key[0],
                "公司": _text(first.get("公司")),
                "公司统一编码": _text(first.get("公司统一编码")),
                "原始公司名称": _text(first.get("原始公司名称")),
                "标准公司名称": _text(first.get("标准公司名称")),
                "公司类型": _text(first.get("公司类型")),
                "同业分类": _text(first.get("同业分类")),
                "报告类型": _text(first.get("报告类型")),
                "报告年度": _text(first.get("报告年度")),
                "报告季度": _text(first.get("报告季度")),
                "报告期": key[1],
                "披露日期": _text(first.get("披露日期")),
                "来源文件": _text(first.get("来源文件")),
                "导入批次": _text(first.get("导入批次")),
            }

    for item in company_periods or ():
        company = _text(item.get("公司") or item.get("公司统一编码"))
        key = _company_key(company)
        report_period = _text(item.get("报告期"))
        if not report_period:
            # 批处理元数据偶尔没有报告期；沿用同一公司在标准数据中已有的报告期，
            # 避免同一家公司被拆成“有报告期”和“空报告期”两组。
            report_period = next(
                (
                    pair["报告期"]
                    for pair in pairs.values()
                    if pair["公司键"] == key and pair["报告期"]
                ),
                "",
            )
        pair_key = (key, report_period)
        base = pairs.get(pair_key)
        if base is None:
            base = {
                "公司键": key,
                "公司": company,
                "公司统一编码": _text(item.get("公司统一编码")),
                "原始公司名称": _text(item.get("原始公司名称") or company),
                "标准公司名称": _text(item.get("标准公司名称") or company),
                "公司类型": _text(item.get("公司类型")),
                "同业分类": _text(item.get("同业分类")),
                "报告类型": _text(item.get("报告类型")),
                "报告年度": _text(item.get("报告年度")),
                "报告季度": _text(item.get("报告季度")),
                "报告期": report_period,
                "披露日期": _text(item.get("披露日期")),
                "来源文件": _text(item.get("来源文件")),
                "导入批次": _text(item.get("导入批次")),
            }
        else:
            for field_name, value in item.items():
                if field_name in base and not base[field_name] and value:
                    base[field_name] = _text(value)
        year, quarter = _report_parts(report_period)
        if not base["报告年度"]:
            base["报告年度"] = year
        if not base["报告季度"]:
            base["报告季度"] = quarter
        pairs[pair_key] = base
    return list(pairs.values())


def build_step7_metric_gap(
    data: pd.DataFrame | None,
    *,
    taxonomy: pd.DataFrame | None = None,
    company_periods: Sequence[dict] | None = None,
    requirement_catalog: pd.DataFrame | None = None,
    require_expected_period: bool = False,
) -> Step7MetricGapReport:
    """Compare STEP3 standard data with the STEP7 requirement catalog."""
    catalog = (
        requirement_catalog.copy()
        if isinstance(requirement_catalog, pd.DataFrame)
        else step7_metric_requirement_catalog(taxonomy)
    )
    if catalog.empty:
        return Step7MetricGapReport(
            missing=pd.DataFrame(columns=MISSING_COLUMNS),
            coverage=pd.DataFrame(columns=_COVERAGE_COLUMNS),
            derived_gaps=pd.DataFrame(columns=_DERIVED_GAP_COLUMNS),
            requirement_catalog=catalog,
            warnings=("STEP7 指标需求目录为空，无法检查缺失指标。",),
        )

    catalog = catalog.copy()
    catalog["指标编码"] = catalog["指标编码"].fillna("").astype(str).str.strip()
    catalog["补录方式"] = catalog["补录方式"].fillna("").astype(str).str.strip()
    catalog["是否派生"] = catalog["是否派生"].fillna("").astype(str).str.strip()
    catalog = catalog[catalog["指标编码"].ne("")].drop_duplicates(
        subset="指标编码",
        keep="first",
    )
    lookup = catalog.set_index("指标编码", drop=False)

    frame = _prepare_data(data)
    pairs = _pair_records(frame, company_periods)
    if not pairs:
        return Step7MetricGapReport(
            missing=pd.DataFrame(columns=MISSING_COLUMNS),
            coverage=pd.DataFrame(columns=_COVERAGE_COLUMNS),
            derived_gaps=pd.DataFrame(columns=_DERIVED_GAP_COLUMNS),
            requirement_catalog=catalog,
            warnings=("没有可检查的公司/报告期组合。",),
        )

    present: dict[tuple[str, str, str], pd.Series] = {}
    present_scopes: dict[tuple[str, str, str], set[str]] = {}
    if not frame.empty:
        numeric = frame[frame["_数值"].notna()]
        for _, row in numeric.iterrows():
            key = (
                str(row["_公司键"]),
                str(row["报告期"]),
                str(row["指标编码"]),
            )
            present.setdefault(key, row)
            present_scopes.setdefault(key, set()).add(_period_group(row["期间口径"]))

    base_codes = tuple(
        catalog.loc[catalog["补录方式"].eq("直接补录"), "指标编码"].tolist()
    )
    derived_codes = tuple(
        catalog.loc[catalog["是否派生"].eq("是"), "指标编码"].tolist()
    )
    # catalog 只带“维护说明/计算逻辑”，依赖关系从派生指标定义中读取；
    # 这里用 DERIVED_METRICS 的 dependencies 作为更可靠来源。
    from .solvency_metric_registry import DERIVED_METRICS

    dependency_map = {
        definition.code: tuple(definition.dependencies)
        for definition in DERIVED_METRICS
    }

    missing_rows: list[dict] = []
    coverage_rows: list[dict] = []
    derived_rows: list[dict] = []

    for pair in pairs:
        company_key = pair["公司键"]
        report_period = pair["报告期"]
        present_codes = {
            code
            for (key, period, code) in present
            if key == company_key and period == report_period
        }

        pair_missing_codes: list[str] = []
        for code in base_codes:
            meta = lookup.loc[code]
            expected_period = _text(meta.get("建议期间口径")) or DEFAULT_AMOUNT_PERIOD
            period_missing = False
            if code in present_codes:
                if require_expected_period and expected_period:
                    scopes = present_scopes.get(
                        (company_key, report_period, code),
                        set(),
                    )
                    if _period_group(expected_period) not in scopes:
                        period_missing = True
                    else:
                        continue
                else:
                    continue
            pair_missing_codes.append(code)
            missing_rows.append(
                {
                    "公司统一编码": pair["公司统一编码"],
                    "公司": pair["公司"],
                    "原始公司名称": pair["原始公司名称"],
                    "标准公司名称": pair["标准公司名称"],
                    "公司类型": pair["公司类型"],
                    "同业分类": pair["同业分类"],
                    "报告类型": pair["报告类型"],
                    "报告年度": pair["报告年度"],
                    "报告季度": pair["报告季度"],
                    "报告期": report_period,
                    "指标编码": code,
                    "指标名称": _text(meta.get("指标名称")),
                    "一级模块": _text(meta.get("一级模块")),
                    "二级模块": _text(meta.get("二级模块")),
                    "数据类型": _text(meta.get("数据类型")),
                    "单位": _text(meta.get("标准单位")),
                    "期间口径": expected_period,
                    "数值": None,
                    "补录优先级": (
                        "高"
                        if code
                        in {
                            "CORE_SOLVENCY_RATIO",
                            "COMBINED_SOLVENCY_RATIO",
                            "ACTUAL_CAPITAL",
                            "MINIMUM_CAPITAL",
                            "SIGNED_PREMIUM",
                            "NET_PROFIT",
                            "COMBINED_RATIO",
                        }
                        else "中"
                    ),
                    "备注": (
                        f"STEP7 需要“{expected_period}”口径，但现有数据只有其他期间口径；"
                        "请补充该期间的数值。"
                        if period_missing
                        else "STEP7 必需的基础指标未在 STEP3 标准数据中找到；"
                        "请按“单位”列口径填写数值。"
                    ),
                }
            )

        pair_derived_gaps: list[dict] = []
        for code in derived_codes:
            if code in present_codes:
                continue
            dependencies = dependency_map.get(code, ())
            missing_dependencies = [
                dependency for dependency in dependencies if dependency not in present_codes
            ]
            if not missing_dependencies:
                # 依赖齐全，STEP5 会自动补算，不计为缺口。
                continue
            meta = lookup.loc[code]
            pair_derived_gaps.append(
                {
                    "公司": pair["公司"],
                    "公司统一编码": pair["公司统一编码"],
                    "报告期": report_period,
                    "派生指标编码": code,
                    "派生指标名称": _text(meta.get("指标名称")),
                    "缺失依赖": "、".join(missing_dependencies),
                    "说明": "依赖指标缺失，派生指标暂时无法计算；补录基础指标后会自动补算。",
                }
            )
        derived_rows.extend(pair_derived_gaps)

        required_count = len(base_codes)
        present_count = required_count - len(pair_missing_codes)
        coverage_rows.append(
            {
                "公司": pair["公司"],
                "公司统一编码": pair["公司统一编码"],
                "同业分类": pair["同业分类"],
                "报告期": report_period,
                "必需基础指标数": required_count,
                "已提取基础指标数": present_count,
                "缺失基础指标数": len(pair_missing_codes),
                "覆盖率": (
                    round(present_count / required_count, 4)
                    if required_count
                    else 1.0
                ),
                "派生指标待计算数": len(pair_derived_gaps),
                "缺失指标编码": "、".join(pair_missing_codes),
            }
        )

    missing = pd.DataFrame(missing_rows, columns=MISSING_COLUMNS)
    if not missing.empty:
        priority_order = {"高": 0, "中": 1, "低": 2}
        missing["_优先级"] = missing["补录优先级"].map(priority_order).fillna(9)
        missing = (
            missing.sort_values(
                ["公司", "_优先级", "指标编码"],
                kind="stable",
            )
            .drop(columns="_优先级")
            .reset_index(drop=True)
        )
    coverage = pd.DataFrame(coverage_rows, columns=_COVERAGE_COLUMNS)
    derived_gaps = pd.DataFrame(derived_rows, columns=_DERIVED_GAP_COLUMNS)

    warnings: list[str] = []
    unknown_codes = sorted(
        catalog.loc[
            catalog["指标名称"].fillna("").astype(str).str.strip().eq("")
            | catalog["指标名称"].fillna("").astype(str).str.strip().eq(
                catalog["指标编码"].fillna("").astype(str).str.strip()
            ),
            "指标编码",
        ].astype(str)
    )
    if unknown_codes:
        warnings.append(
            "以下 STEP7 指标编码未在指标字典中匹配到名称/单位：" + "、".join(unknown_codes)
        )
    return Step7MetricGapReport(
        missing=missing,
        coverage=coverage,
        derived_gaps=derived_gaps,
        requirement_catalog=catalog,
        warnings=tuple(warnings),
    )


def _style_sheet(sheet, *, freeze: str = "A2", header_fill: str = "00338D") -> None:
    header_fill_obj = PatternFill("solid", fgColor=header_fill)
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10)
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet.freeze_panes = freeze
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = header_fill_obj
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    for column_cells in sheet.iter_cols(
        min_row=1,
        max_row=min(sheet.max_row, 300),
    ):
        values = []
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            values.append(sum(2 if ord(char) > 127 else 1 for char in value))
        width = min(max(max(values, default=8) + 3, 10), 46)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def _instructions_frame(report: Step7MetricGapReport) -> pd.DataFrame:
    return pd.DataFrame(
        [
            (
                "用途",
                "本表由 STEP3 标准化完成后自动生成，用于补录 STEP7 公司报告需要、"
                "但本次提取没有取得的指标。",
            ),
            (
                "怎么填",
                f"只需在“{MISSING_SHEET_NAME}”工作表的“数值”列填写数字；"
                "其他列（公司、报告期、指标编码、期间口径、单位）已按标准口径预填，请勿改动。",
            ),
            (
                "单位",
                "金额类标准单位是万元；如果原始数据是元，请把“单位”改为“元”，"
                "合并时会自动除以 10000。百分比直接填 130.98 表示 130.98%。",
            ),
            (
                "期间口径",
                "默认按指标类型预填（存量指标本季度末数、流量指标本年累计数）。"
                "如实际披露口径不同，可在“期间口径”列修改后合并。",
            ),
            (
                "合并键",
                "公司统一编码 + 报告期 + 指标编码 + 期间口径。"
                "公司统一编码为空时使用公司名称。",
            ),
            (
                "合并方式",
                "保存本文件后，回到 STEP3 页面，在同一检查区上传本工作簿，"
                "点击“应用补录并合并”，系统会追加数值并自动补算派生指标。",
            ),
            (
                "本次缺口",
                f"公司/报告期组合 {report.company_count} 个；"
                f"缺失基础指标组合 {report.missing_count} 条；"
                f"涉及指标 {report.missing_code_count} 个；"
                f"当前基础指标覆盖率 {report.coverage_ratio:.1%}。",
            ),
            (
                "派生指标",
                f"另有 {len(report.derived_gaps)} 条派生指标因依赖缺失暂时无法计算，"
                f"详见“{DERIVED_GAP_SHEET_NAME}”工作表；补录基础指标后会自动补算。",
            ),
        ],
        columns=["项目", "说明"],
    )


def missing_metric_template_workbook_bytes(
    report: Step7MetricGapReport,
) -> bytes:
    """Build the fillable missing-metric workbook."""
    fill = report.missing.copy()
    for column in MISSING_COLUMNS:
        if column not in fill.columns:
            fill[column] = ""
    fill = fill.reindex(columns=MISSING_COLUMNS)
    catalog = report.requirement_catalog.copy()
    catalog_columns = [
        column
        for column in (
            "指标编码",
            "指标名称",
            "一级模块",
            "二级模块",
            "标准单位",
            "数据类型",
            "补录方式",
            "建议期间口径",
            "Step7使用图表",
        )
        if column in catalog.columns
    ]

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _instructions_frame(report).to_excel(
            writer,
            sheet_name=INSTRUCTION_SHEET_NAME,
            index=False,
        )
        fill.to_excel(writer, sheet_name=MISSING_SHEET_NAME, index=False)
        report.coverage.to_excel(
            writer,
            sheet_name=COVERAGE_SHEET_NAME,
            index=False,
        )
        catalog[catalog_columns].to_excel(
            writer,
            sheet_name=REQUIREMENT_SHEET_NAME,
            index=False,
        )
        if not report.derived_gaps.empty:
            report.derived_gaps.to_excel(
                writer,
                sheet_name=DERIVED_GAP_SHEET_NAME,
                index=False,
            )

        for name in writer.book.sheetnames:
            sheet = writer.book[name]
            if name == INSTRUCTION_SHEET_NAME:
                _style_sheet(sheet)
                sheet.column_dimensions["A"].width = 18
                sheet.column_dimensions["B"].width = 100
            else:
                _style_sheet(sheet)

        sheet = writer.book[MISSING_SHEET_NAME]
        value_col = MISSING_COLUMNS.index("数值") + 1
        value_letter = get_column_letter(value_col)
        max_row = max(sheet.max_row, 2)
        # 数值列用浅黄色提示“这里需要填写”。
        fill_obj = PatternFill("solid", fgColor="FFF2CC")
        for row in range(2, max_row + 1):
            cell = sheet.cell(row=row, column=value_col)
            cell.fill = fill_obj
            cell.number_format = "#,##0.00####"
            cell.alignment = Alignment(horizontal="right", vertical="top")
        sheet.cell(row=1, column=value_col).comment = Comment(
            "请按“单位”列口径填写数值。金额标准单位万元；"
            "如原始数据为元，请把单位改为“元”，合并时自动换算。",
            "STEP3",
        )
        number_validation = DataValidation(
            type="decimal",
            operator="between",
            formula1="-1000000000000000",
            formula2="1000000000000000",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="数值格式",
            error="请填写数字，不要带千分位或单位文字。",
        )
        sheet.add_data_validation(number_validation)
        number_validation.add(f"{value_letter}2:{value_letter}{max_row}")

        unit_col = MISSING_COLUMNS.index("单位") + 1
        unit_validation = DataValidation(
            type="list",
            formula1='"元,千元,万元,百元,亿元,十亿元,%,倍"',
            allow_blank=True,
        )
        sheet.add_data_validation(unit_validation)
        unit_validation.add(
            f"{get_column_letter(unit_col)}2:"
            f"{get_column_letter(unit_col)}{max_row}"
        )

        period_col = MISSING_COLUMNS.index("期间口径") + 1
        period_validation = DataValidation(
            type="list",
            formula1='"' + ",".join(_PERIOD_SUGGESTIONS) + '"',
            allow_blank=True,
        )
        sheet.add_data_validation(period_validation)
        period_validation.add(
            f"{get_column_letter(period_col)}2:"
            f"{get_column_letter(period_col)}{max_row}"
        )
        # 前 10 列保持在视野内，横向滚动填写数值时仍能看到公司和指标。
        sheet.freeze_panes = "K2"
    return output.getvalue()


def _find_header_row(preview: pd.DataFrame) -> int | None:
    for index in range(min(20, len(preview))):
        values = {str(value).strip() for value in preview.iloc[index].tolist()}
        if "指标编码" in values:
            return index
    return None


def read_missing_metric_fills(workbook_bytes: bytes) -> MissingFillParseResult:
    """Read a filled missing-metric workbook and validate numeric values."""
    excel = pd.ExcelFile(io.BytesIO(workbook_bytes))
    sheet_name = (
        MISSING_SHEET_NAME
        if MISSING_SHEET_NAME in excel.sheet_names
        else next(
            (
                name
                for name in excel.sheet_names
                if "缺失" in str(name) or "补录" in str(name)
            ),
            excel.sheet_names[0],
        )
    )
    preview = pd.read_excel(excel, sheet_name=sheet_name, header=None, nrows=20)
    header_row = _find_header_row(preview)
    if header_row is None:
        raise ValueError(f"补录表 {sheet_name} 未找到“指标编码”表头。")
    frame = pd.read_excel(excel, sheet_name=sheet_name, header=header_row).fillna("")
    frame.columns = [str(column).strip() for column in frame.columns]

    required_columns = {"公司", "报告期", "指标编码", "期间口径", "数值"}
    missing_columns = sorted(required_columns - set(frame.columns))
    if missing_columns:
        raise ValueError("补录表缺少必要列：" + "、".join(missing_columns))

    rows: list[dict] = []
    skipped: list[dict] = []
    for index, row in frame.iterrows():
        excel_row = int(index) + header_row + 2
        value = _to_number(row.get("数值"))
        raw_value = _text(row.get("数值"))
        company = _text(row.get("公司"))
        report_period = _text(row.get("报告期"))
        code = _text(row.get("指标编码"))
        period = _text(row.get("期间口径"))
        # 未填写的模板行直接忽略，不计入“跳过”清单；
        # 只有写了内容但无法解析成数字时才提示。
        if value is None:
            if not raw_value:
                continue
            skipped.append(
                {
                    "Excel行号": excel_row,
                    "公司": company,
                    "指标编码": code,
                    "原因": f"数值“{raw_value}”无法解析为数字",
                }
            )
            continue
        if not company or not report_period or not code:
            skipped.append(
                {
                    "Excel行号": excel_row,
                    "公司": company,
                    "指标编码": code,
                    "原因": "公司、报告期或指标编码为空",
                }
            )
            continue
        item = row.to_dict()
        item["数值"] = value
        item["_Excel行号"] = excel_row
        rows.append(item)

    return MissingFillParseResult(
        rows=pd.DataFrame(rows),
        skipped=pd.DataFrame(skipped),
        warnings=(),
    )


def _catalog_lookup(
    taxonomy: pd.DataFrame | None,
    catalog: pd.DataFrame | None,
) -> dict[str, dict]:
    frame = (
        catalog.copy()
        if isinstance(catalog, pd.DataFrame)
        else step7_metric_requirement_catalog(taxonomy)
    )
    lookup: dict[str, dict] = {}
    for _, row in frame.iterrows():
        code = _text(row.get("指标编码"))
        if code:
            lookup[code] = row.to_dict()
    return lookup


def merge_missing_metric_fills(
    base_data: pd.DataFrame | None,
    fills: pd.DataFrame,
    *,
    taxonomy: pd.DataFrame | None = None,
    requirement_catalog: pd.DataFrame | None = None,
) -> MissingFillMergeResult:
    """Append filled missing metrics to the STEP3 narrow table and recompute derived."""
    if not isinstance(fills, pd.DataFrame) or fills.empty:
        return MissingFillMergeResult(
            data=_prepare_data(base_data).drop(
                columns=["_公司键", "_数值"],
                errors="ignore",
            ),
            added_count=0,
            skipped_count=0,
            warnings=("补录表中没有可合并的数值。",),
        )

    base = _prepare_data(base_data)
    lookup = _catalog_lookup(taxonomy, requirement_catalog)
    existing_keys = {
        (
            str(row["_公司键"]),
            str(row["报告期"]),
            str(row["指标编码"]),
            str(row["期间口径"]),
        )
        for _, row in base.iterrows()
        if pd.notna(row["_数值"])
    }

    new_rows: list[dict] = []
    warnings: list[str] = []
    skipped = 0
    for _, row in fills.iterrows():
        code = _text(row.get("指标编码"))
        company = _text(row.get("公司"))
        report_period = _text(row.get("报告期"))
        period = _text(row.get("期间口径")) or DEFAULT_AMOUNT_PERIOD
        value = _to_number(row.get("数值"))
        if not code or not company or not report_period or value is None:
            skipped += 1
            continue
        company_key = _company_key(company)
        base_match = base[base["_公司键"].eq(company_key)]
        if not base_match.empty:
            base_row = base_match.iloc[0]
            row = row.copy()
            for column in (
                "公司",
                "原始公司名称",
                "标准公司名称",
                "公司统一编码",
                "公司类型",
                "同业分类",
                "报告类型",
                "报告年度",
                "报告季度",
                "报告期",
                "披露日期",
                "来源文件",
                "导入批次",
            ):
                if not _text(row.get(column)):
                    row[column] = base_row.get(column, "")
            company = _text(row.get("公司")) or company
        duplicate_key = (company_key, report_period, code, period)
        if duplicate_key in existing_keys:
            skipped += 1
            warnings.append(
                f"{company} {report_period} {code}（{period}）已存在数值，已跳过补录。"
            )
            continue

        meta = lookup.get(code, {})
        standard_unit = _text(meta.get("标准单位")) or _text(row.get("单位"))
        data_type = _text(meta.get("数据类型")) or _text(row.get("数据类型"))
        input_unit = _text(row.get("单位")) or standard_unit
        conversion_note = ""
        if (
            data_type == "金额"
            and input_unit in _CURRENCY_UNITS
            and standard_unit in _CURRENCY_UNITS
            and input_unit != standard_unit
        ):
            converted, note = convert_amount_to_target_unit(
                value,
                input_unit,
                standard_unit,
                metric_code=code,
            )
            if isinstance(converted, (int, float)):
                value = float(converted)
                conversion_note = note
            elif note:
                warnings.append(f"{company} {code} 单位未换算：{note}")

        year, quarter = _report_parts(report_period)
        note_parts = [
            "STEP3 缺失指标人工补录",
            f"原填写单位：{input_unit or '未填写'}",
        ]
        if conversion_note:
            note_parts.append(conversion_note)
        if _text(row.get("备注")):
            note_parts.append(_text(row.get("备注")))

        record = {column: "" for column in STANDARD_COLUMNS}
        record.update(
            {
                "公司": company,
                "原始公司名称": _text(row.get("原始公司名称")) or company,
                "标准公司名称": _text(row.get("标准公司名称")) or company,
                "公司统一编码": _text(row.get("公司统一编码")) or f"AUTO_{company_key}",
                "公司类型": _text(row.get("公司类型")),
                "同业分类": _text(row.get("同业分类")),
                "报告类型": _text(row.get("报告类型")),
                "报告年度": _text(row.get("报告年度")) or year,
                "报告季度": _text(row.get("报告季度")) or quarter,
                "报告期": report_period,
                "披露日期": _text(row.get("披露日期")),
                "一级模块": _text(meta.get("一级模块")) or _text(row.get("一级模块")),
                "二级模块": _text(meta.get("二级模块")) or _text(row.get("二级模块")),
                "指标编码": code,
                "指标名称": _text(meta.get("指标名称")) or _text(row.get("指标名称")),
                "期间口径": period,
                "数值": value,
                "单位": standard_unit or input_unit,
                "数据类型": data_type or _text(row.get("数据类型")),
                "是否预测": "否",
                "来源页码": _text(row.get("来源页码")),
                "原始披露值": "",
                "备注": "；".join(part for part in note_parts if part),
                "来源类型": "人工补录",
                "指标属性": _text(meta.get("指标属性")) or _text(row.get("指标属性")) or "披露",
                "来源文件": _text(row.get("来源文件")),
                "来源工作表": "缺失指标补录",
                "导入批次": _text(row.get("导入批次")),
                "计算逻辑": _text(meta.get("计算逻辑")),
            }
        )
        new_rows.append(record)
        existing_keys.add(duplicate_key)

    merged = pd.concat(
        [
            base.drop(columns=["_公司键", "_数值"], errors="ignore"),
            pd.DataFrame(new_rows, columns=STANDARD_COLUMNS),
        ],
        ignore_index=True,
    )
    if new_rows:
        try:
            merged = add_missing_derived_metrics(merged)
        except Exception as exc:  # pragma: no cover - defensive fallback
            warnings.append(f"派生指标自动补算失败，已保留基础补录结果：{exc}")

    return MissingFillMergeResult(
        data=merged.reindex(columns=STANDARD_COLUMNS).reset_index(drop=True),
        added_count=len(new_rows),
        skipped_count=skipped,
        warnings=tuple(warnings),
    )
