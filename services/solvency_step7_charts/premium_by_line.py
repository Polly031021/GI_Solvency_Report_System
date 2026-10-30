"""分险种保费收入：数据读取、同比增速与年度汇总表。

数据优先从根目录《偿付能力行业集成数据_公司名称与分类规范化.xlsx》的
“分险种保费收入”工作表读取；若该工作表不存在，则回退到根目录
《分险种保费收入.xlsx》，最后回退到本模块内置的同一份数据。
"""

from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import altair as alt
import pandas as pd

from ..report_i18n import t
from .primitives import apply_industry_chart_theme

ROOT = Path(__file__).resolve().parents[2]
PREMIUM_BY_LINE_PATH = ROOT / "分险种保费收入.xlsx"
INTEGRATED_WORKBOOK_PATH = ROOT / "偿付能力行业集成数据_公司名称与分类规范化.xlsx"
INTEGRATED_SHEET_NAME = "分险种保费收入"
DISPLAY_UNIT = "亿元"
MOTOR_LINE = "机动车辆保险"
AMOUNT_COLUMN = "原保险保费收入"
GROWTH_COLUMN = "同比增速"
BASE_COLUMN = "基期年度"

SOURCE_INTEGRATED = "integrated"
SOURCE_STANDALONE = "standalone"
SOURCE_EMBEDDED = "embedded"

LINE_NAMES_EN: dict[str, str] = {
    "机动车辆保险": "Motor vehicle insurance",
    "责任险": "Liability insurance",
    "农业保险": "Agricultural insurance",
    "健康险": "Health insurance",
    "意外险": "Accident insurance",
    "企业财产保险": "Commercial property insurance",
    "家庭财产保险": "Home property insurance",
    "工程保险": "Engineering insurance",
    "保证保险": "Guarantee insurance",
}

# Stable colours keep the same line of business recognisable in the legend.
LINE_COLORS: dict[str, str] = {
    "机动车辆保险": "#00338D",
    "健康险": "#FD349C",
    "农业保险": "#00C0AE",
    "责任险": "#7213EA",
    "意外险": "#F1C44D",
    "企业财产保险": "#1E49E2",
    "保证保险": "#ED2124",
    "家庭财产保险": "#B497FF",
    "工程保险": "#63EBB2",
}
FALLBACK_LINE_COLOR = "#94A3B8"
# 同比增速折线使用与柱体不同的深橙色，避免同色导致柱线走势难以区分。
GROWTH_LINE_COLOR = "#C2410C"
MISSING_MARK_COLOR = "#94A3B8"


def _readable_text_color(background: str) -> str:
    """Return a legible label colour for a filled bar."""
    value = str(background or "").strip().lstrip("#")
    if len(value) != 6:
        return "#FFFFFF"
    try:
        channels = [int(value[index : index + 2], 16) / 255 for index in (0, 2, 4)]
    except ValueError:
        return "#FFFFFF"
    linear = [
        channel / 12.92 if channel <= 0.03928 else ((channel + 0.055) / 1.055) ** 2.4
        for channel in channels
    ]
    luminance = 0.2126 * linear[0] + 0.7152 * linear[1] + 0.0722 * linear[2]
    return "#0C233C" if luminance > 0.45 else "#FFFFFF"


DISCLOSED = "已披露"

NOT_DISCLOSED = "未披露"

# 展示口径：只列示五个主要险种，并且只展示最近四个年度。
DISPLAY_LINES: tuple[str, ...] = (
    "机动车辆保险",
    "健康险",
    "农业保险",
    "责任险",
    "意外险",
)
DISPLAY_YEARS = 4

# 与根目录《分险种保费收入.xlsx》完全一致的内置副本（单位：亿元）。
# 仅在源文件缺失时使用。
EMBEDDED_PREMIUM_BY_LINE: tuple[tuple[int, str, float], ...] = (
    (2025, "机动车辆保险", 9409.0),
    (2025, "责任险", 1428.0),
    (2025, "农业保险", 1555.0),
    (2025, "健康险", 2274.0),
    (2025, "意外险", 593.0),
    (2024, "机动车辆保险", 9137.0),
    (2024, "责任险", 1372.0),
    (2024, "农业保险", 1484.0),
    (2024, "健康险", 2043.0),
    (2024, "意外险", 534.0),
    (2023, "机动车辆保险", 8673.0),
    (2023, "责任险", 1268.0),
    (2023, "农业保险", 1430.0),
    (2023, "健康险", 1752.0),
    (2023, "意外险", 509.0),
    (2022, "企业财产保险", 553.0),
    (2022, "家庭财产保险", 164.0),
    (2022, "机动车辆保险", 8210.0),
    (2022, "工程保险", 145.0),
    (2022, "责任险", 1148.0),
    (2022, "保证保险", 552.0),
    (2022, "农业保险", 1219.0),
    (2022, "健康险", 1580.0),
    (2022, "意外险", 574.0),
    (2021, "企业财产保险", 520.0),
    (2021, "家庭财产保险", 98.0),
    (2021, "机动车辆保险", 7773.0),
    (2021, "工程保险", 144.0),
    (2021, "责任险", 1018.0),
    (2021, "保证保险", 521.0),
    (2021, "农业保险", 976.0),
    (2021, "健康险", 1378.0),
    (2021, "意外险", 627.0),
    (2020, "企业财产保险", 490.0),
    (2020, "家庭财产保险", 91.0),
    (2020, "机动车辆保险", 8245.0),
    (2020, "工程保险", 138.0),
    (2020, "责任险", 901.0),
    (2020, "保证保险", 689.0),
    (2020, "农业保险", 815.0),
    (2020, "健康险", 1114.0),
    (2020, "意外险", 541.0),
    (2019, "企业财产保险", 464.0),
    (2019, "家庭财产保险", 91.0),
    (2019, "机动车辆保险", 8188.0),
    (2019, "工程保险", 118.0),
    (2019, "责任险", 753.0),
    (2019, "保证保险", 844.0),
    (2019, "农业保险", 672.0),
    (2019, "健康险", 840.0),
    (2019, "意外险", 527.0),
)

_BASE_COLUMNS = ["年度", "险种", AMOUNT_COLUMN]
_SOURCE_TEXT = {
    SOURCE_INTEGRATED: (
        "根目录《偿付能力行业集成数据_公司名称与分类规范化.xlsx》"
        "的“分险种保费收入”工作表",
        "the premium-by-line sheet in the integrated solvency workbook",
    ),
    SOURCE_STANDALONE: (
        "根目录《分险种保费收入.xlsx》",
        "the premium-by-line workbook in the project root",
    ),
    SOURCE_EMBEDDED: ("内置数据", "built-in data"),
}


@dataclass(frozen=True)
class PremiumByLineOverview:
    """Latest-year aggregates used by the module insight card."""

    year: int
    total: float
    previous_total: float | None
    growth: float | None
    top_line: str
    top_share: float
    fastest_line: str
    fastest_growth: float
    slowest_line: str
    slowest_growth: float
    active_lines: tuple[str, ...]
    legacy_lines: tuple[str, ...]


def line_display_name(name: object) -> str:
    """Return the Chinese or English display name of one line of business."""
    text = str(name or "").strip()
    return t(text, LINE_NAMES_EN.get(text, text))


def premium_by_line_source_key() -> str:
    """Return which workbook supplied the data (integrated/standalone/embedded)."""
    return _load_with_source()[1]


def premium_by_line_source_text() -> str:
    """Return the localised source description used in captions/footnotes."""
    zh, en = _SOURCE_TEXT.get(premium_by_line_source_key(), _SOURCE_TEXT[SOURCE_EMBEDDED])
    return t(zh, en)


def _amount_factor(text: str) -> float:
    """Convert the workbook's declared amount unit into 亿元."""
    if "十亿元" in text:
        return 10.0
    if "百万元" in text:
        return 0.01
    if "亿元" in text:
        return 1.0
    if "万元" in text:
        return 0.0001
    if "元" in text:
        return 1e-8
    return 1.0


def _parse_tabular(raw: pd.DataFrame) -> pd.DataFrame:
    """Parse a 年份 / 险种 / 保费收入 table without relying on fixed columns."""
    empty = pd.DataFrame(columns=_BASE_COLUMNS)
    if not isinstance(raw, pd.DataFrame) or raw.empty:
        return empty
    header_index: int | None = None
    for index in range(min(12, len(raw))):
        values = [
            str(value).strip()
            for value in raw.iloc[index].tolist()
            if not pd.isna(value)
        ]
        has_year = any(("年份" in value or "年度" in value) for value in values)
        has_line = any("险种" in value for value in values)
        if has_year and has_line:
            header_index = index
            break
    if header_index is None:
        return empty

    header = [str(value).strip() for value in raw.iloc[header_index].tolist()]
    year_column = next(
        (index for index, name in enumerate(header) if "年份" in name or "年度" in name),
        None,
    )
    line_column = next(
        (index for index, name in enumerate(header) if "险种" in name),
        None,
    )
    amount_column = next(
        (
            index
            for index, name in enumerate(header)
            if ("保费" in name or "金额" in name) and "同比" not in name
        ),
        None,
    )
    unit_column = next(
        (index for index, name in enumerate(header) if name.strip() == "单位"),
        None,
    )
    if year_column is None or line_column is None:
        return empty
    if amount_column is None:
        amount_column = 2
    default_factor = _amount_factor(" ".join(header))

    rows: list[dict[str, object]] = []
    for _, row in raw.iloc[header_index + 1:].iterrows():
        year = pd.to_numeric(row.iloc[year_column], errors="coerce")
        line = row.iloc[line_column]
        amount = pd.to_numeric(row.iloc[amount_column], errors="coerce")
        line_text = "" if pd.isna(line) else str(line).strip()
        if pd.isna(year) or not line_text or pd.isna(amount):
            continue
        factor = default_factor
        if unit_column is not None:
            unit_text = row.iloc[unit_column]
            if not pd.isna(unit_text) and str(unit_text).strip():
                factor = _amount_factor(str(unit_text))
        rows.append(
            {
                "年度": int(year),
                "险种": line_text,
                AMOUNT_COLUMN: float(amount) * factor,
            }
        )
    if not rows:
        return empty
    return (
        pd.DataFrame(rows, columns=_BASE_COLUMNS)
        .drop_duplicates(["年度", "险种"], keep="last")
        .reset_index(drop=True)
    )


def _read_integrated_sheet() -> pd.DataFrame:
    """Read the by-line sheet from the large integrated workbook quickly."""
    if not INTEGRATED_WORKBOOK_PATH.exists():
        return pd.DataFrame(columns=_BASE_COLUMNS)
    try:
        # read_only avoids parsing the 60k+ row sheets just to read this table.
        from openpyxl import load_workbook

        workbook = load_workbook(
            INTEGRATED_WORKBOOK_PATH,
            read_only=True,
            data_only=True,
        )
        try:
            if INTEGRATED_SHEET_NAME not in workbook.sheetnames:
                return pd.DataFrame(columns=_BASE_COLUMNS)
            values = list(workbook[INTEGRATED_SHEET_NAME].iter_rows(values_only=True))
        finally:
            workbook.close()
        return _parse_tabular(pd.DataFrame(values))
    except Exception:
        return pd.DataFrame(columns=_BASE_COLUMNS)


def _read_standalone_workbook() -> pd.DataFrame:
    if not PREMIUM_BY_LINE_PATH.exists():
        return pd.DataFrame(columns=_BASE_COLUMNS)
    try:
        return _parse_tabular(pd.read_excel(PREMIUM_BY_LINE_PATH, header=None))
    except Exception:
        return pd.DataFrame(columns=_BASE_COLUMNS)


def _attach_growth(frame: pd.DataFrame) -> pd.DataFrame:
    """Sort each line by year and calculate YoY growth on available years."""
    result = frame.copy()
    result[AMOUNT_COLUMN] = pd.to_numeric(result[AMOUNT_COLUMN], errors="coerce")
    result["年度"] = pd.to_numeric(result["年度"], errors="coerce")
    result = result.dropna(subset=["年度", AMOUNT_COLUMN, "险种"]).copy()
    if result.empty:
        return pd.DataFrame(columns=[*_BASE_COLUMNS, GROWTH_COLUMN, BASE_COLUMN])
    result["年度"] = result["年度"].astype(int)
    result = result.sort_values(["险种", "年度"], kind="mergesort").reset_index(drop=True)

    growth: list[float] = [float("nan")] * len(result)
    bases: list[str] = [""] * len(result)
    for indices in result.groupby("险种", sort=False).groups.values():
        previous_year: int | None = None
        previous_value: float | None = None
        for index in indices:
            value = float(result.at[index, AMOUNT_COLUMN])
            if previous_value is not None and previous_value != 0:
                growth[index] = (value / previous_value - 1.0) * 100.0
                bases[index] = str(previous_year)
            previous_year = int(result.at[index, "年度"])
            previous_value = value
    result[GROWTH_COLUMN] = growth
    result[BASE_COLUMN] = bases
    return result


@lru_cache(maxsize=1)
def _load_with_source() -> tuple[pd.DataFrame, str]:
    """Load the frame once, preferring the integrated workbook's new sheet."""
    for source, loader in (
        (SOURCE_INTEGRATED, _read_integrated_sheet),
        (SOURCE_STANDALONE, _read_standalone_workbook),
    ):
        frame = loader()
        if not frame.empty:
            return _attach_growth(frame), source
    embedded = pd.DataFrame(EMBEDDED_PREMIUM_BY_LINE, columns=_BASE_COLUMNS)
    return _attach_growth(embedded), SOURCE_EMBEDDED


def load_premium_by_line_data() -> pd.DataFrame:
    """Load and normalise the by-line premium data (cached per process)."""
    return _load_with_source()[0]


def ordered_premium_lines(frame: pd.DataFrame) -> list[str]:
    """Order lines by latest reported amount, active lines before retired ones."""
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        return []
    latest = (
        frame.sort_values("年度")
        .groupby("险种", sort=False, as_index=False)
        .tail(1)
    )
    max_year = int(frame["年度"].max())
    active = latest[latest["年度"].eq(max_year)].sort_values(
        AMOUNT_COLUMN, ascending=False
    )
    legacy = latest[latest["年度"].lt(max_year)].sort_values(
        AMOUNT_COLUMN, ascending=False
    )
    return [*active["险种"].astype(str), *legacy["险种"].astype(str)]


def premium_line_frame(frame: pd.DataFrame, line_name: str) -> pd.DataFrame:
    """Return one line of business ordered by year."""
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        return pd.DataFrame(columns=[*_BASE_COLUMNS, GROWTH_COLUMN, BASE_COLUMN])
    selected = frame[frame["险种"].astype(str).eq(str(line_name))].copy()
    return selected.sort_values("年度").reset_index(drop=True)


def premium_by_line_window_years(
    frame: pd.DataFrame,
    count: int = DISPLAY_YEARS,
) -> list[int]:
    """Return the latest count years available in the frame."""
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        return []
    years = sorted(int(year) for year in frame["年度"].dropna().unique())
    return years[-count:] if count > 0 else years


def premium_by_line_display_frame(
    frame: pd.DataFrame,
    *,
    lines: tuple[str, ...] = DISPLAY_LINES,
    years: list[int] | None = None,
) -> pd.DataFrame:
    """Restrict the source frame to the presented lines and latest years.

    Growth is calculated on the full history before this filter, so the first
    visible year still carries its year-on-year growth and base year.
    """
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        return frame
    if years is None:
        years = premium_by_line_window_years(frame)
    selected = frame[
        frame["险种"].astype(str).isin(lines)
        & frame["年度"].astype(int).isin(years)
    ].copy()
    return selected.sort_values(
        ["年度", "险种"], kind="mergesort"
    ).reset_index(drop=True)


def premium_by_line_chart_frame(
    frame: pd.DataFrame,
    *,
    line_name: str | None = None,
    years: list[int] | None = None,
) -> pd.DataFrame:
    """Return long rows for the per-line charts, including missing slots.

    A missing line-year combination keeps its place on the year axis and is
    marked with an explicit “×”. Pass line_name to build a single line frame;
    omit it to build every line.
    """
    columns = ["年度", "险种", AMOUNT_COLUMN, GROWTH_COLUMN, BASE_COLUMN, "披露状态"]
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        return pd.DataFrame(columns=columns)
    if years is None:
        years = sorted(int(year) for year in frame["年度"].dropna().unique())
    else:
        years = sorted(int(year) for year in years)
    lines = [str(line_name)] if line_name is not None else ordered_premium_lines(frame)
    if line_name is not None and str(line_name) not in set(frame["险种"].astype(str)):
        lines = []
    rows: list[dict[str, object]] = []
    for line_name in lines:
        line_frame = premium_line_frame(frame, line_name)
        by_year = {
            int(row.年度): row
            for row in line_frame.itertuples(index=False)
        }
        for year in years:
            record = by_year.get(year)
            if record is None:
                rows.append(
                    {
                        "年度": str(year),
                        "险种": line_name,
                        AMOUNT_COLUMN: float("nan"),
                        GROWTH_COLUMN: float("nan"),
                        BASE_COLUMN: "",
                        "披露状态": NOT_DISCLOSED,
                    }
                )
                continue
            growth = getattr(record, GROWTH_COLUMN)
            rows.append(
                {
                    "年度": str(year),
                    "险种": line_name,
                    AMOUNT_COLUMN: float(getattr(record, AMOUNT_COLUMN)),
                    GROWTH_COLUMN: float(growth) if pd.notna(growth) else float("nan"),
                    BASE_COLUMN: str(getattr(record, BASE_COLUMN) or ""),
                    "披露状态": DISCLOSED,
                }
            )
    return pd.DataFrame(rows, columns=columns)


def _premium_year_axis(years: list[str]) -> alt.Axis:
    return alt.Axis(
        values=years,
        labels=True,
        labelAngle=0,
        labelFontSize=10,
        labelPadding=6,
        domain=True,
        ticks=True,
        tickSize=4,
        grid=False,
    )


def _premium_year_encoding(years: list[str]) -> alt.X:
    return alt.X(
        "年度:O",
        sort=years,
        title=None,
        scale=alt.Scale(paddingInner=0.22, paddingOuter=0.3),
        axis=_premium_year_axis(years),
    )


def _premium_amount_tooltip(amount_title: str) -> list[alt.Tooltip]:
    return [
        alt.Tooltip("年度:N", title=t("年度", "Year")),
        alt.Tooltip("险种:N", title=t("险种", "Line of business")),
        alt.Tooltip(f"{AMOUNT_COLUMN}:Q", title=amount_title, format=",.1f"),
        alt.Tooltip(
            f"{GROWTH_COLUMN}:Q",
            title=t("同比增速", "YoY growth"),
            format="+.1f",
        ),
        alt.Tooltip(f"{BASE_COLUMN}:N", title=t("基期年度", "Base year")),
        alt.Tooltip("披露状态:N", title=t("披露状态", "Status")),
    ]


def build_premium_by_line_combo_chart(
    frame: pd.DataFrame,
    line_name: str,
    *,
    years: list[int] | None = None,
    width: int | str = "container",
    height: int = 300,
) -> alt.Chart:
    """One line of business: premium bars and YoY growth in one chart.

    The bars use the left axis (RMB 100 mn) and the growth line uses an
    independent right axis (%), so a single chart carries both series.
    """
    chart_frame = premium_by_line_chart_frame(
        frame, line_name=line_name, years=years
    )
    if chart_frame.empty:
        raise ValueError(f"没有可绘制的险种数据：{line_name}")
    year_labels = chart_frame["年度"].astype(str).tolist()
    disclosed = chart_frame[chart_frame["披露状态"].eq(DISCLOSED)].copy()
    missing = chart_frame[chart_frame["披露状态"].eq(NOT_DISCLOSED)].copy()
    growth = disclosed.dropna(subset=[GROWTH_COLUMN]).copy()
    if growth.empty:
        raise ValueError(f"没有可用的同比增速数据：{line_name}")
    growth["增速标签"] = growth[GROWTH_COLUMN].map(
        lambda value: f"{float(value):+.1f}%"
    )
    color = LINE_COLORS.get(str(line_name), FALLBACK_LINE_COLOR)
    growth_color = GROWTH_LINE_COLOR
    amount_title = t(
        "原保险保费收入（亿元）",
        "Original premium income (RMB 100 mn)",
    )
    growth_title = t("同比增速（%）", "YoY growth (%)")
    max_value = float(
        pd.to_numeric(disclosed[AMOUNT_COLUMN], errors="coerce").max() or 0.0
    )
    amount_scale = alt.Scale(
        domain=[0, max_value * 1.26 if max_value > 0 else 1],
        zero=True,
        nice=False,
    )
    values = pd.to_numeric(growth[GROWTH_COLUMN], errors="coerce")
    lower = min(0.0, float(values.min()))
    upper = max(0.0, float(values.max()))
    pad = max(4.0, (upper - lower) * 0.2)
    growth_scale = alt.Scale(
        domain=[lower - pad, upper + pad],
        zero=False,
        nice=False,
    )
    # Keep growth annotations in a clear band above the bars and line.
    growth["标签位置"] = upper + pad * 0.45
    bars = (
        alt.Chart(disclosed)
        .mark_bar(
            size=34,
            color=color,
            cornerRadiusTopLeft=2,
            cornerRadiusTopRight=2,
        )
        .encode(
            x=_premium_year_encoding(year_labels),
            y=alt.Y(
                f"{AMOUNT_COLUMN}:Q",
                title=amount_title,
                scale=amount_scale,
                axis=alt.Axis(
                    orient="left",
                    format=",.0f",
                    tickCount=5,
                    domain=True,
                    ticks=True,
                    tickSize=4,
                    grid=True,
                ),
            ),
            tooltip=_premium_amount_tooltip(amount_title),
        )
    )
    # Place amount labels at the bar base so they never collide with the growth
    # labels, which sit above the line points near the top of the plot.
    amount_labels = (
        alt.Chart(disclosed.assign(金额标签位置=0.0))
        .mark_text(
            dy=-8,
            baseline="bottom",
            fontSize=9,
            fontWeight=600,
            color=_readable_text_color(color),
        )
        .encode(
            x=alt.X("年度:O", sort=year_labels, axis=None),
            y=alt.Y("金额标签位置:Q", scale=amount_scale, axis=None),
            text=alt.Text(f"{AMOUNT_COLUMN}:Q", format=",.0f"),
        )
    )
    zero_rule = (
        alt.Chart(pd.DataFrame({GROWTH_COLUMN: [0.0]}))
        .mark_rule(color="#94A3B8", strokeDash=[4, 4], strokeWidth=1)
        .encode(y=alt.Y(f"{GROWTH_COLUMN}:Q", scale=growth_scale, axis=None))
    )
    growth_line = (
        alt.Chart(growth)
        .mark_line(
            strokeWidth=2.1,
            color=growth_color,
            point=alt.OverlayMarkDef(
                size=52,
                filled=True,
                color=growth_color,
                stroke="#FFFFFF",
                strokeWidth=1.2,
            ),
        )
        .encode(
            x=alt.X("年度:O", sort=year_labels, axis=None),
            y=alt.Y(
                f"{GROWTH_COLUMN}:Q",
                title=growth_title,
                scale=growth_scale,
                axis=alt.Axis(
                    format=".0f",
                    labelExpr="datum.label + '%'",
                    orient="right",
                    tickCount=5,
                    domain=True,
                    ticks=True,
                    tickSize=4,
                    grid=False,
                ),
            ),
            tooltip=[
                alt.Tooltip("年度:N", title=t("年度", "Year")),
                alt.Tooltip("险种:N", title=t("险种", "Line of business")),
                alt.Tooltip(
                    f"{AMOUNT_COLUMN}:Q",
                    title=amount_title,
                    format=",.1f",
                ),
                alt.Tooltip(
                    f"{GROWTH_COLUMN}:Q",
                    title=growth_title,
                    format="+.1f",
                ),
                alt.Tooltip(f"{BASE_COLUMN}:N", title=t("基期年度", "Base year")),
            ],
        )
    )
    growth_labels = (
        alt.Chart(growth)
        .mark_text(
            fontSize=9,
            fontWeight=600,
            color="#0C233C",
        )
        .encode(
            x=alt.X("年度:O", sort=year_labels, axis=None),
            y=alt.Y("标签位置:Q", scale=growth_scale, axis=None),
            text=alt.Text("增速标签:N"),
        )
    )
    layers: list[alt.Chart] = [
        bars,
        amount_labels,
        zero_rule,
        growth_line,
        growth_labels,
    ]
    if not missing.empty:
        missing = missing.assign(标记位置=0.0)
        layers.append(
            alt.Chart(missing)
            .mark_text(
                text="×",
                dy=-6,
                fontSize=12,
                fontWeight="bold",
                color=MISSING_MARK_COLOR,
            )
            .encode(
                x=alt.X("年度:O", sort=year_labels, axis=None),
                y=alt.Y("标记位置:Q", scale=growth_scale, axis=None),
                tooltip=[
                    alt.Tooltip("年度:N", title=t("年度", "Year")),
                    alt.Tooltip("险种:N", title=t("险种", "Line of business")),
                    alt.Tooltip("披露状态:N", title=t("披露状态", "Status")),
                ],
            )
        )
    chart = (
        alt.layer(*layers)
        .resolve_scale(y="independent")
        .properties(width=width)
    )
    return apply_industry_chart_theme(chart, height=height).properties(
        padding={"top": 10, "right": 12, "bottom": 8, "left": 4}
    )

def prepare_premium_by_line_overview(
    frame: pd.DataFrame,
) -> PremiumByLineOverview | None:
    """Aggregate the latest-year active lines for the module insight card."""
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        return None
    latest_rows = (
        frame.sort_values("年度")
        .groupby("险种", sort=False, as_index=False)
        .tail(1)
    )
    max_year = int(frame["年度"].max())
    active = latest_rows[latest_rows["年度"].eq(max_year)].copy()
    if active.empty:
        return None
    active = active.sort_values(AMOUNT_COLUMN, ascending=False)
    total = float(pd.to_numeric(active[AMOUNT_COLUMN], errors="coerce").sum())
    active_names = tuple(active["险种"].astype(str))
    top_row = active.iloc[0]
    top_line = str(top_row["险种"])
    top_value = float(top_row[AMOUNT_COLUMN])
    top_share = (top_value / total * 100.0) if total > 0 else 0.0

    previous_year = max_year - 1
    previous = frame[
        frame["年度"].eq(previous_year) & frame["险种"].astype(str).isin(active_names)
    ]
    previous_names = set(previous["险种"].astype(str))
    if previous_names == set(active_names):
        previous_total = float(
            pd.to_numeric(previous[AMOUNT_COLUMN], errors="coerce").sum()
        )
    else:
        previous_total = None
    growth = (
        (total / previous_total - 1.0) * 100.0
        if previous_total not in (None, 0) and total > 0
        else None
    )

    growth_rows = active.dropna(subset=[GROWTH_COLUMN]).copy()
    if growth_rows.empty:
        fastest_line = top_line
        fastest_growth = 0.0
        slowest_line = top_line
        slowest_growth = 0.0
    else:
        fastest = growth_rows.loc[growth_rows[GROWTH_COLUMN].idxmax()]
        slowest = growth_rows.loc[growth_rows[GROWTH_COLUMN].idxmin()]
        fastest_line = str(fastest["险种"])
        fastest_growth = float(fastest[GROWTH_COLUMN])
        slowest_line = str(slowest["险种"])
        slowest_growth = float(slowest[GROWTH_COLUMN])

    legacy_lines = tuple(
        latest_rows.loc[latest_rows["年度"].lt(max_year), "险种"].astype(str)
    )
    return PremiumByLineOverview(
        year=max_year,
        total=total,
        previous_total=previous_total,
        growth=growth,
        top_line=top_line,
        top_share=top_share,
        fastest_line=fastest_line,
        fastest_growth=fastest_growth,
        slowest_line=slowest_line,
        slowest_growth=slowest_growth,
        active_lines=active_names,
        legacy_lines=legacy_lines,
    )
